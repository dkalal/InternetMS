"""Purchase-workspace spreadsheet parsing and tenant-scoped batch validation.

This module never writes business records. The save boundary re-runs validation.
"""

import csv
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO

from django.core.exceptions import ValidationError
from django.core.validators import DecimalValidator
from django.db.models import Q
from django.db.models.functions import Lower, Upper
from openpyxl import load_workbook

from products.models import Product, ProductCategory, UnitOfMeasure

MAX_ROWS = 200
MAX_TEXT = 100_000
MAX_FILE = 250_000
MAX_SOURCE_ROWS = 10_000
HEADERS = (
    'sku', 'product_name', 'quantity', 'unit_cost', 'category', 'unit',
    'selling_price', 'brand', 'model_number', 'serialized', 'track_expiry',
    'batch_reference', 'expiry_date', 'serial_numbers', 'staging_reference',
)
LEGACY_HEADERS = ('sku', 'quantity', 'unit_cost', 'batch_reference', 'expiry_date', 'serial_numbers')
DECIMAL_FIELDS = {
    'quantity': (16, 6, True), 'unit_cost': (16, 6, False),
    'selling_price': (10, 2, True),
}
LENGTHS = {'sku': 64, 'product_name': 200, 'category': 120, 'unit': 50,
           'brand': 120, 'model_number': 120, 'batch_reference': 100,
           'serial_numbers': 10000, 'staging_reference': 64}


class ImportInputError(ValueError):
    pass


class ImportColumnMappingRequired(ImportInputError):
    def __init__(self, message, structure):
        super().__init__(message)
        self.structure = structure


def _as_text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


HEADER_ALIASES = {
    'product': 'product_name', 'name': 'product_name', 'qty': 'quantity',
    'cost': 'unit_cost', 'purchase_price': 'unit_cost', 'price': 'unit_cost',
    'batch': 'batch_reference', 'expiry': 'expiry_date', 'serials': 'serial_numbers',
    'model': 'model_number', 'sales_unit': 'unit', 'is_serialized': 'serialized',
}
STRUCTURE_HELP = ('We could not identify the spreadsheet columns. Copy the selected cells '
                  'directly from Excel or upload the template.')
FIELD_LABELS = {
    'sku': 'SKU', 'product_name': 'Product name', 'quantity': 'Quantity',
    'unit_cost': 'Unit cost (TZS)', 'category': 'Category', 'unit': 'Sales / stock unit',
    'selling_price': 'Selling price (TZS)', 'brand': 'Brand', 'model_number': 'Model',
    'serialized': 'Track serial numbers', 'track_expiry': 'Track expiry dates',
    'batch_reference': 'Batch reference', 'expiry_date': 'Expiry date',
    'serial_numbers': 'Serial numbers', 'staging_reference': 'Shared product reference',
}


def _header(value):
    key = re.sub(r'[\s-]+', '_', value.strip().lower().lstrip('\ufeff'))
    return HEADER_ALIASES.get(key, key)


def _numeric_text(value):
    """Only English grouped thousands with a decimal dot are unambiguous here."""
    if re.fullmatch(r'[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?', value):
        return value.replace(',', '')
    return value


def _read_input(*, text='', upload=None):
    """Read cells once, retaining empty columns and physical source positions."""
    if upload:
        if upload.size > MAX_FILE:
            raise ImportInputError('XLSX file exceeds 250 KB.')
        if not upload.name.lower().endswith('.xlsx'):
            raise ImportInputError('Upload an .xlsx template. Macro-enabled files are not accepted.')
        book = None
        try:
            book = load_workbook(BytesIO(upload.read()), read_only=True, data_only=False)
            if 'Purchase rows' not in book.sheetnames:
                raise ImportInputError('The workbook needs a Purchase rows sheet.')
            sheet = book['Purchase rows']
            if sheet.max_column > len(HEADERS) or sheet.max_row > MAX_SOURCE_ROWS:
                raise ImportInputError('The worksheet supports at most 15 columns and 10,000 source rows, including blank rows.')
            raw = []
            for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                row = [_as_text(cell) for cell in cells]
                if any(row):
                    raw.append((number, row))
                    if len(raw) > MAX_ROWS + 1:
                        raise ImportInputError('The workbook exceeds 200 non-empty purchase rows.')
        except ImportInputError:
            raise
        except Exception as exc:
            raise ImportInputError('The XLSX file could not be read.') from exc
        finally:
            if book is not None:
                book.close()
    else:
        if len(text) > MAX_TEXT:
            raise ImportInputError('Pasted text exceeds 100,000 characters.')
        try:
            reader = csv.reader(StringIO(text, newline=''), delimiter='\t', strict=True)
            raw = []
            start = 1
            for cells in reader:
                if any(cell.strip() for cell in cells):
                    raw.append((start, cells))
                start = reader.line_num + 1
        except csv.Error as exc:
            raise ImportInputError('Check spreadsheet quoting in the pasted rows.') from exc
    if not raw:
        raise ImportInputError('Enter at least one purchase row.')
    return raw


def _has_header(raw):
    return sum(_header(cell) in HEADERS for cell in raw[0][1]) >= 2


def _describe_columns(raw, *, has_header=None):
    if has_header is not None and not isinstance(has_header, bool):
        raise ImportInputError('Specify whether the first row contains column headings.')
    has_header = _has_header(raw) if has_header is None else has_header
    data = raw[1:] if has_header else raw
    if not 1 <= len(data) <= MAX_ROWS:
        raise ImportInputError(f'Enter between 1 and {MAX_ROWS} non-empty purchase rows.')
    width = max(len(cells) for _, cells in raw)
    if not 3 <= width <= len(HEADERS):
        raise ImportInputError(STRUCTURE_HELP + ' Expected 3 to 15 tab-separated columns.')
    headers = [_header(cell) for cell in raw[0][1]] if has_header else []
    columns = []
    for index in range(width):
        heading = headers[index] if index < len(headers) else ''
        examples = []
        for _, cells in data[:20]:
            value = cells[index] if index < len(cells) else ''
            value = _as_text(value)[:100] or '(blank)'
            if value not in examples:
                examples.append(value)
            if len(examples) == 3:
                break
        columns.append({
            'source_column': index + 1,
            'heading': (_as_text(raw[0][1][index])[:100] if has_header and index < len(raw[0][1]) else ''),
            'examples': examples,
            # Suggest by known header only. Never infer financial fields from values.
            'field': heading if heading in HEADERS and headers.count(heading) == 1 else '',
        })
    return {'row_count': len(data), 'column_count': width, 'has_header': has_header,
            'columns': columns, 'fields': [{'value': key, 'label': FIELD_LABELS[key]} for key in HEADERS]}


def inspect_input(*, text='', upload=None, has_header=None):
    """Bounded read-only structure inspection, shared by XLSX and TSV."""
    return _describe_columns(_read_input(text=text, upload=upload), has_header=has_header)


def _mapping_required(message, raw, *, has_header=None):
    return ImportColumnMappingRequired(message, _describe_columns(raw, has_header=has_header))


def _validated_mapping(mapping, raw, *, has_header):
    width = max(len(cells) for _, cells in raw)
    if not isinstance(has_header, bool):
        raise ImportInputError('Confirm whether the first row contains column headings.')
    if (not isinstance(mapping, list) or not 3 <= len(mapping) <= len(HEADERS)
            or not all(isinstance(key, str) for key in mapping)):
        raise ImportInputError('Column mapping must contain 3 to 15 field selections.')
    if len(mapping) != width:
        raise ImportInputError(f'Column mapping expects {width} selections; received {len(mapping)}. Choose columns again.')
    seen = set()
    for position, key in enumerate(mapping, 1):
        if key == '__ignore__':
            continue
        if key not in HEADERS:
            raise ImportInputError(f'Column {position}: choose a purchase field or explicitly select Do not import.')
        if key in seen:
            raise ImportInputError(f'{FIELD_LABELS[key]} is assigned more than once. Choose a different field for column {position}.')
        seen.add(key)
    if not {'quantity', 'unit_cost'}.issubset(seen) or not {'sku', 'product_name'}.intersection(seen):
        raise ImportInputError('Map Quantity, Unit cost, and SKU or Product name before review.')
    return tuple(mapping)


def parse_input(*, text='', upload=None, input_format='auto', column_mapping=None, has_header=None):
    if input_format not in ('auto', 'mixed', 'legacy'):
        raise ImportInputError('Choose a supported spreadsheet format.')
    raw = _read_input(text=text, upload=upload)
    source_raw = raw
    first = tuple(_header(cell) for cell in raw[0][1])
    # A header is recognizable by several names, not by one SKU coincidentally named SKU.
    is_header = _has_header(raw)
    if column_mapping is not None:
        keys = _validated_mapping(column_mapping, raw, has_header=has_header)
        raw = raw[1:] if has_header else raw
    elif is_header:
        if len(first) != len(set(first)):
            raise _mapping_required('Duplicate or ambiguous column headers. Use each column name once.', raw)
        unknown = [key for key in first if key not in HEADERS]
        if unknown:
            raise _mapping_required('Unknown column headers: ' + ', '.join(unknown) + '. Choose columns or use the template headers.', raw)
        if not {'quantity', 'unit_cost'}.issubset(first) or not {'sku', 'product_name'}.intersection(first):
            raise _mapping_required('Headers need Quantity, Unit cost, and SKU or Product name.', raw)
        keys = first
        raw = raw[1:]
    elif upload:
        raise ImportInputError('The Purchase rows sheet needs recognized column headers. Choose columns or download the template.')
    elif input_format == 'mixed':
        keys = HEADERS
    elif input_format == 'legacy':
        keys = LEGACY_HEADERS
    else:
        # Headerless legacy is safe only when the first two numeric columns fit.
        probe = raw[0][1]
        if len(probe) > 6:
            raise _mapping_required(f'Row {raw[0][0]}: expected at most 6 legacy columns, received {len(probe)}. '
                                    'Choose columns to identify this spreadsheet layout.', raw)
        if len(probe) == 1:
            probe = probe[0].split()
        if 3 <= len(probe) <= 6 and all(_looks_numeric(v) or v.startswith('=') for v in probe[1:3]):
            keys = LEGACY_HEADERS
        else:
            raise ImportInputError(STRUCTURE_HELP + ' For headerless rows, select the Mixed or Legacy format.')
    if not raw or len(raw) > MAX_ROWS:
        raise ImportInputError(f'Enter between 1 and {MAX_ROWS} non-empty purchase rows.')
    result = []
    for index, cells in raw:
        if len(cells) == 1 and keys != LEGACY_HEADERS:
            raise ImportInputError(f'Row {index}: ' + STRUCTURE_HELP)
        if len(cells) == 1 and keys == LEGACY_HEADERS:
            # Compatibility: only an unambiguous SKU + numeric quantity/cost,
            # optionally followed by single-token batch/date/serial values.
            parts = cells[0].split()
            if 3 <= len(parts) <= 6 and all(_looks_numeric(v) or v.startswith('=') for v in parts[1:3]):
                cells = parts
            else:
                raise ImportInputError(f'Row {index}: ' + STRUCTURE_HELP)
        if len(cells) > len(keys):
            raise ImportInputError(f'Row {index}: expected at most {len(keys)} columns, received {len(cells)}. ' + STRUCTURE_HELP)
        row = {key: _as_text(cells[pos]) if pos < len(cells) else ''
               for pos, key in enumerate(keys) if key != '__ignore__'}
        for key in ('quantity', 'unit_cost', 'selling_price'):
            if key in row:
                row[key] = _numeric_text(row[key])
        for key in ('quantity', 'unit_cost'):
            value = row.get(key, '')
            if value and not _looks_numeric(value) and not value.startswith('='):
                reason = (f'Row {index}, {key.replace("_", " ")}: expected a number. '
                          f'Source column {keys.index(key) + 1} is mapped to {FIELD_LABELS[key]}. '
                          'Choose columns if this differs from your layout; '
                          'use 1250.50 or 1,250.50, without currency symbols.')
                if column_mapping is None:
                    raise _mapping_required(reason, source_raw)
                raise ImportInputError(reason)
        row['row'] = index
        row['input_errors'] = []
        for key, max_length in LENGTHS.items():
            if len(row.get(key, '')) > max_length:
                row['input_errors'].append(f'{key.replace("_", " ").title()} is too long.')
        if any(value.startswith('=') for key, value in row.items()
               if key not in ('row', 'input_errors') and isinstance(value, str)):
            row['input_errors'].append('Formula cells are not accepted. Paste values as text.')
        result.append(row)
    return result


def _looks_numeric(value):
    try:
        Decimal(_numeric_text(value))
        return True
    except (InvalidOperation, ValueError):
        return False


def _decimal(value, key, errors):
    digits, places, positive = DECIMAL_FIELDS[key]
    label = key.replace('_', ' ').capitalize()
    if not value:
        errors.append(f'{label} is required.')
        return None
    try:
        number = Decimal(_numeric_text(value))
        if not number.is_finite() or number < 0 or (positive and number <= 0):
            raise InvalidOperation
        DecimalValidator(digits, places)(number)
        return number
    except (InvalidOperation, ValueError, ValidationError):
        errors.append(f'{label} has an invalid value or precision.')
        return None


def _flag(value, label, errors):
    normalized = str(value).strip().lower()
    if normalized in ('', 'no', 'false', '0'):
        return False
    if normalized in ('yes', 'true', '1'):
        return True
    errors.append(f'{label} must be Yes or No.')
    return False


def validate_rows(rows, *, organization, for_save=False):
    """Return display rows; every catalog lookup is scoped to the active tenant."""
    if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_ROWS:
        raise ImportInputError('The batch must contain 1 to 200 rows.')
    sku_set = {str(r.get('sku', '')).strip().upper() for r in rows if r.get('sku')}
    ids = {int(str(r['product_id'])) for r in rows
           if str(r.get('product_id', '')).isdigit() and len(str(r['product_id'])) <= 18}
    catalog = list(Product.objects.filter(tenant=organization).annotate(normalized_sku=Upper('sku')).filter(
        Q(normalized_sku__in=sku_set) | Q(pk__in=ids)
    ).select_related('sales_unit'))
    by_sku = {p.sku.upper(): p for p in catalog}
    by_id = {p.pk: p for p in catalog}
    candidate_names = {str(r.get('product_name', '')).strip().casefold() for r in rows
                       if r.get('product_name') and not r.get('sku')}
    suggestions_by_name = {}
    if candidate_names:
        suggestions = Product.objects.filter(
            tenant=organization, is_active=True, item_type=Product.ItemType.PHYSICAL, track_stock=True
        ).annotate(normalized_name=Lower('name')).filter(
            normalized_name__in=candidate_names
        ).select_related('sales_unit').order_by('name', 'pk')[:600]
        for candidate in suggestions:
            matches = suggestions_by_name.setdefault(candidate.name.casefold(), [])
            if len(matches) < 10:
                matches.append(candidate)
    categories = {c.name.casefold(): c for c in ProductCategory.objects.filter(tenant=organization, is_active=True).prefetch_related('allowed_units')}
    allowed_units = {c.pk: {unit.pk for unit in c.allowed_units.all()} for c in categories.values()}
    units = {u.name.casefold(): u for u in UnitOfMeasure.objects.filter(tenant=organization, is_active=True)}
    seen_new_sku = {}
    definitions = {}
    seen_serials = set()
    results = []
    for number, supplied in enumerate(rows, 1):
        row = {key: _as_text(supplied.get(key, '')) for key in HEADERS}
        for key in DECIMAL_FIELDS:
            row[key] = _numeric_text(row[key])
        source_row = supplied.get('row', number)
        row['row'] = source_row if isinstance(source_row, int) and 0 < source_row <= MAX_TEXT else number
        errors = list(supplied.get('input_errors', []))
        for key, length in LENGTHS.items():
            if len(row[key]) > length:
                errors.append(f'{key.replace("_", " ").title()} is too long.')
        if any(row[key].startswith('=') for key in HEADERS):
            errors.append('Formula cells are not accepted.')
        row['sku'] = row['sku'].upper()
        excluded = str(supplied.get('excluded', '')).lower() in ('true', '1', 'yes')
        resolution = str(supplied.get('resolution', '')).strip()
        row['resolution'] = resolution
        selected_id = str(supplied.get('product_id', '')).strip()
        known = by_sku.get(row['sku']) if row['sku'] else None
        selected = by_id.get(int(selected_id)) if selected_id.isdigit() and len(selected_id) <= 18 else None
        if selected and supplied.get('expected_sku') and selected.sku.upper() != str(supplied['expected_sku']).upper():
            errors.append('The selected catalog SKU changed after preview. Search and select the product again.')
        if selected and supplied.get('expected_serialized') is not None and str(supplied['expected_serialized']).lower() in ('true', 'false'):
            if selected.is_serialized != (str(supplied['expected_serialized']).lower() == 'true'):
                errors.append('The selected product tracking settings changed after preview.')
        if selected and supplied.get('expected_expiry') is not None and str(supplied['expected_expiry']).lower() in ('true', 'false'):
            if selected.track_expiry != (str(supplied['expected_expiry']).lower() == 'true'):
                errors.append('The selected product expiry setting changed after preview.')
        if selected and supplied.get('expected_unit_id'):
            expected_unit = str(supplied['expected_unit_id'])
            if not expected_unit.isdigit() or selected.sales_unit_id != int(expected_unit):
                errors.append('The selected product unit changed after preview.')
        product = None
        status = 'unresolved'
        if excluded:
            status = 'excluded'
        elif known and not (known.is_active and known.item_type == Product.ItemType.PHYSICAL and known.track_stock):
            errors.append('SKU belongs to an inactive, service, or non-stock product.')
        elif known and resolution != 'new':
            if selected_id and selected_id != str(known.pk):
                errors.append('Selected product conflicts with the supplied SKU.')
            else:
                product = known
        elif resolution == 'existing':
            if not selected or not (selected.is_active and selected.item_type == Product.ItemType.PHYSICAL and selected.track_stock):
                errors.append('Choose an eligible active stock product from this catalog.')
            else:
                product = selected
        elif resolution == 'new':
            status = 'new'
            if not row['product_name']:
                errors.append('Product name is required for a new product.')
            category = categories.get(row['category'].casefold())
            unit = units.get(row['unit'].casefold())
            if not category:
                errors.append('Choose an active category in this business.')
            if not unit:
                errors.append('Choose an active unit in this business.')
            if category and unit and unit.pk not in allowed_units[category.pk]:
                errors.append('The unit is not allowed for this category.')
            row['selling_price_value'] = _decimal(row['selling_price'], 'selling_price', errors)
            row['serialized_value'] = _flag(row['serialized'], 'Serialized', errors)
            row['track_expiry_value'] = _flag(row['track_expiry'], 'Track expiry', errors)
            if row['sku']:
                from products.sku import validate_manual_sku
                from django.core.exceptions import ValidationError
                try:
                    validate_manual_sku(row['sku'])
                except ValidationError as exc:
                    errors.extend(exc.messages)
                if known:
                    errors.append('This SKU already belongs to a catalog product.')
                prior = seen_new_sku.setdefault(row['sku'], number)
                if prior != number and row['staging_reference'] != results[prior - 1].get('staging_reference'):
                    errors.append(f'New-product SKU is already proposed on row {prior}.')
            reference = row['staging_reference'] or f'row-{number}'
            definition = tuple(row.get(k) for k in ('sku', 'product_name', 'category', 'unit', 'selling_price', 'brand', 'model_number', 'serialized', 'track_expiry'))
            if reference in definitions and definitions[reference] != definition:
                errors.append('Rows sharing the same proposed product have conflicting product definitions.')
            definitions[reference] = definition
            row['staging_reference'] = reference
            row['category_id'] = category.pk if category else None
            row['unit_id'] = unit.pk if unit else None
        else:
            if row['sku'] and not known:
                errors.append('No active stock product matches this SKU. Choose an existing product or Create new.')
            elif not row['sku']:
                errors.append('SKU is missing. Choose an existing product or Create new.')
        if product:
            status = 'matched'
            row['product'] = {'id': product.pk, 'name': product.name, 'sku': product.sku,
                              'brand': product.brand, 'model_number': product.model_number,
                              'text': str(product), 'unit': product.sales_unit.label if product.sales_unit_id else product.measure_unit,
                              'unit_id': product.sales_unit_id,
                              'serialized': product.is_serialized, 'expiry': product.track_expiry}
            row['product_id'] = product.pk
            row['expected_sku'] = product.sku
            row['expected_serialized'] = product.is_serialized
            row['expected_expiry'] = product.track_expiry
            row['expected_unit_id'] = product.sales_unit_id
        else:
            row['product'] = None
            row['product_id'] = selected.pk if selected else None
        row['suggestions'] = [
            {'id': candidate.pk, 'sku': candidate.sku, 'name': candidate.name,
             'brand': candidate.brand, 'model_number': candidate.model_number,
             'unit': candidate.sales_unit.label if candidate.sales_unit_id else candidate.measure_unit,
             'unit_id': candidate.sales_unit_id, 'serialized': candidate.is_serialized,
             'expiry': candidate.track_expiry}
            for candidate in suggestions_by_name.get(row['product_name'].casefold(), [])
            if (not row['brand'] or candidate.brand.casefold() == row['brand'].casefold())
            and (not row['model_number'] or candidate.model_number.casefold() == row['model_number'].casefold())
            and (not row['unit'] or candidate.sales_unit and candidate.sales_unit.name.casefold() == row['unit'].casefold())
        ][:3]
        quantity = _decimal(row['quantity'], 'quantity', errors)
        cost = _decimal(row['unit_cost'], 'unit_cost', errors)
        if row['expiry_date']:
            try:
                date.fromisoformat(row['expiry_date'])
            except ValueError:
                errors.append('Expiry date must use YYYY-MM-DD.')
        serials = [s.strip() for s in row['serial_numbers'].replace('|', ',').replace('\n', ',').split(',') if s.strip()]
        for serial in ([] if excluded else serials):
            normalized = serial.upper()
            if normalized in seen_serials:
                errors.append(f'Serial {serial} occurs in another row.')
            seen_serials.add(normalized)
        serialized = product.is_serialized if product else row.get('serialized_value', False)
        if serialized and quantity is not None and (quantity != quantity.to_integral_value() or len(serials) != int(quantity)):
            errors.append('Serialized products need one unique serial per whole unit.')
        if (product or status == 'new') and not serialized and serials:
            errors.append('Quantity-based products must not have serial numbers.')
        expiry_tracked = product.track_expiry if product else row.get('track_expiry_value', False)
        if (product or status == 'new') and expiry_tracked and not row['expiry_date']:
            errors.append('An expiry date is required for this product.')
        if not excluded and status == 'unresolved' and not errors:
            errors.append('Choose an existing product or Create new.') if for_save else None
        row['serial_numbers'] = '\n'.join(serials)
        row['quantity'] = format(quantity, 'f') if quantity is not None else row['quantity']
        row['unit_cost'] = format(cost, 'f') if cost is not None else row['unit_cost']
        row['excluded'] = excluded
        row['status'] = 'excluded' if excluded else ('unresolved' if status == 'unresolved' and len(errors) == 1
                                                    else 'invalid' if errors else status)
        row['valid'] = excluded or (not errors and status in ('matched', 'new'))
        row['errors'] = [] if excluded else errors
        results.append(row)
    from .serial_validation import existing_serials
    existing = existing_serials(organization=organization, serials=[
        serial for result in results if not result['excluded']
        for serial in result['serial_numbers'].splitlines()
    ])
    for result in results:
        conflicts = sorted({serial.upper() for serial in result['serial_numbers'].splitlines()} & existing)
        if conflicts and not result['excluded']:
            result['errors'].append('Serials already exist in this organization: ' + ', '.join(conflicts[:10]))
            result['valid'] = False
            result['status'] = 'invalid'
    return results


def submitted_rows(value):
    if not value:
        return []
    if len(value) > MAX_TEXT:
        raise ImportInputError('Applied batch exceeds 100,000 characters.')
    try:
        rows = json.loads(value)
    except (TypeError, ValueError) as exc:
        raise ImportInputError('Applied batch data is invalid. Review the import again.') from exc
    if not isinstance(rows, list):
        raise ImportInputError('Applied batch data is invalid.')
    if not all(isinstance(row, dict) for row in rows):
        raise ImportInputError('Applied batch rows are invalid.')
    return rows
