from __future__ import annotations

import hashlib
import json
import uuid
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.validators import DecimalValidator
from django.db import IntegrityError, transaction
from django.db.models import Count, Exists, F, Max, OuterRef, Prefetch, Q, Subquery, Sum
from django.http import Http404, HttpResponse, JsonResponse
from django.core.paginator import Paginator
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from billing.models import BillingDocument
from billing.services import BillingService, BillingServiceError
from internetservices.number_display import format_quantity
from products.models import Product, ProductCategory, UnitOfMeasure
from products.pricing import cost_floor_for
from users.permissions import PermissionCode, has_tenant_permission, require_permission
from users.tenancy import require_organization

from .forms import (
    CartForm,
    CartPaymentForm,
    CartLineDiscountForm,
    CartLineForm,
    InventoryImportForm,
    InventorySettingsForm,
    ProductCategoryForm,
    PurchaseForm,
    PurchaseLinesFormSet,
    QuickProductCategoryForm,
    QuickPurchaseProductForm,
    QuickSupplierForm,
    StockAdjustmentForm,
    SupplierForm,
    SupplierPaymentForm,
)
from .imports import commit_import, template_workbook, validate_workbook, workbook_bytes
from .models import (
    Cart,
    CartLine,
    CartSerialSelection,
    DocumentSerialSelection,
    ImportJob,
    HistoricalInventoryRecord,
    InventoryBalance,
    InventorySale,
    InventorySaleLine,
    InventorySettings,
    Purchase,
    PurchaseImportSubmission,
    PurchaseLine,
    StockMovement,
    StockUnit,
    Supplier,
)
from .reporting import DATED_REPORTS, ReportFilterForm, report_rows, report_totals, filter_report_rows, valuation_bridge, report_display
from .numbering import PurchaseReferenceNumberService
from .services import CartService, InventoryError, InventoryService, audit, money
from .purchase_import import (
    HEADERS as MIXED_HEADERS, ImportColumnMappingRequired, ImportInputError,
    inspect_input, parse_input, submitted_rows, validate_rows,
)
from .cart_lifecycle import business_linked, cart_access, visible_carts, stale
from .purchase_batch import apply_purchase_batch


def _scope(request, permission):
    organization = require_organization(request)
    require_permission(request, permission)
    return organization


def _cart_workspace(cart):
    """Build the same pricing/totals data for HTML and POS JSON responses."""
    lines = list(cart.lines.select_related('product').prefetch_related('serial_selections__stock_unit'))
    for line in lines:
        _, line.pos_pricing_mode = CartService.line_pricing(
            product=line.product, quantity=line.quantity, customer=cart.customer,
            sale_pricing_category=cart.sale_pricing_category,
        )
    subtotal = sum((line.line_total for line in lines), Decimal('0.00'))
    taxable_subtotal = sum(
        (line.line_total for line in lines if not line.product_id or line.product.tax_eligible),
        Decimal('0.00'),
    )
    _, tax, grand_total = BillingService.compute_totals(
        tax_rate=cart.tax_rate,
        line_items=lines,
        discount_amount=cart.discount_amount,
    )
    return {
        'cart': cart, 'lines': lines, 'subtotal': subtotal,
        'taxable_subtotal': taxable_subtotal, 'tax': tax,
        'grand_total': grand_total,
    }


def _is_pos_request(request):
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest' and 'application/json' in request.headers.get('Accept', '')


def _form_error_summary(form):
    messages_list = []
    for field_name, field_errors in form.errors.items():
        label = form.fields[field_name].label if field_name in form.fields else 'Sale details'
        messages_list.extend(f'{label}: {error}' for error in field_errors)
    return ' '.join(messages_list)


def _pos_response(request, cart, *, message='', level='success', status=200, errors=None):
    workspace = _cart_workspace(cart)
    workspace['can_register_payment'] = has_tenant_permission(
        request.user, cart.tenant, PermissionCode.PAYMENT_REGISTER, membership=request.membership,
    )
    workspace['can_apply_line_discount'] = has_tenant_permission(
        request.user, cart.tenant, PermissionCode.CART_DISCOUNT_APPLY, membership=request.membership,
    )
    payload = {
        'ok': status < 400,
        'message': message,
        'level': level,
        'cart_html': render_to_string('inventory/includes/pos_cart_lines.html', workspace, request=request),
        'checkout_html': render_to_string('inventory/includes/pos_checkout.html', workspace, request=request),
        'subtotal': str(workspace['subtotal']),
        'taxable_subtotal': str(workspace['taxable_subtotal']),
        'discount': str(cart.discount_amount),
        'tax': str(workspace['tax']),
        'tax_rate': str(cart.tax_rate),
        'grand_total': str(workspace['grand_total']),
        'line_count': len(workspace['lines']),
        'lines': [
            {
                'product_id': line.product_id,
                'quantity': str(line.quantity),
                'unit_price': str(line.unit_price),
                'line_total': str(line.line_total),
                'pricing_mode': line.pos_pricing_mode,
            }
            for line in workspace['lines']
        ],
    }
    if errors is not None:
        payload['errors'] = errors
    return JsonResponse(payload, status=status)


@login_required
def dashboard(request):
    organization = _scope(request, PermissionCode.STOCK_VIEW)
    finance_all = has_tenant_permission(
        request.user, organization, PermissionCode.FINANCE_SALES_VIEW_ALL,
        membership=request.membership,
    )
    balances = list(InventoryBalance.objects.filter(tenant=organization).select_related('product'))
    settings_obj, _ = InventorySettings.objects.get_or_create(
        organization=organization, tenant=organization,
        defaults={'walk_in_customer_label': 'Walk-in Customer'},
    )
    cutoff = timezone.now() - timedelta(days=settings_obj.dead_stock_days)
    sold_product_ids = set(StockMovement.objects.filter(
        tenant=organization, movement_type=StockMovement.MovementType.SALE_OUT, created_at__gte=cutoff
    ).values_list('product_id', flat=True))
    fast_cutoff = timezone.now() - timedelta(days=settings_obj.fast_moving_days)
    fast_count = sum(
        1 for row in StockMovement.objects.filter(
            tenant=organization, movement_type=StockMovement.MovementType.SALE_OUT, created_at__gte=fast_cutoff
        ).values('product_id').annotate(units=Sum('quantity'))
        if -row['units'] >= settings_obj.fast_moving_min_units
    )
    stock_products = Product.objects.filter(
        tenant=organization,
        item_type=Product.ItemType.PHYSICAL,
        track_stock=True,
    )
    context = {
        'product_count': Product.objects.filter(tenant=organization).count(),
        'total_stock_units': sum((balance.quantity for balance in balances), Decimal('0.00')),
        'low_stock_count': stock_products.filter(
            quantity__gt=0, quantity__lte=F('reorder_threshold')
        ).count(),
        'out_of_stock_count': stock_products.filter(quantity__lte=0).count(),
        'dead_stock_count': sum(
            1 for balance in balances if balance.quantity > 0 and balance.product_id not in sold_product_ids
        ),
        'fast_moving_count': fast_count,
        'draft_cart_count': Cart.objects.filter(tenant=organization, status=Cart.Status.DRAFT).count(),
    }
    if finance_all:
        context['total_selling_value'] = sum(
            (balance.quantity * balance.product.selling_price for balance in balances), Decimal('0.00')
        )
        context['total_technician_value'] = sum(
            (balance.quantity * balance.product.effective_technician_price for balance in balances), Decimal('0.00')
        )
        context['total_wholesale_value'] = sum(
            (
                balance.quantity * balance.product.price_for_sale_category(
                    sale_pricing_category=Product.PricingMode.WHOLESALE,
                    quantity=balance.quantity,
                )
                for balance in balances
            ),
            Decimal('0.00'),
        )
        context['draft_purchase_count'] = Purchase.objects.filter(tenant=organization, status=Purchase.Status.DRAFT).count()
        context['recent_movements'] = StockMovement.objects.filter(tenant=organization).select_related('product', 'created_by')[:10]
        context['recent_purchases'] = Purchase.objects.filter(tenant=organization).select_related(
            'supplier', 'created_by'
        ).annotate(item_count=Count('lines'))[:5]
    if has_tenant_permission(request.user, organization, PermissionCode.COST_REPORT_VIEW, membership=request.membership):
        context['total_stock_value'] = sum((balance.total_value for balance in balances), Decimal('0.00'))
    return render(request, 'inventory/dashboard.html', context)


@login_required
def category_list(request):
    organization = _scope(request, PermissionCode.PRODUCT_VIEW)
    query = request.GET.get('q', '').strip()
    categories = ProductCategory.objects.filter(tenant=organization).annotate(
        product_count=Count('products'),
    )
    if query:
        categories = categories.filter(Q(name__icontains=query) | Q(description__icontains=query))
    categories = categories.order_by('name')
    category_stats = {
        'total': categories.count(),
        'active': categories.filter(is_active=True).count(),
        'products': sum(category.product_count for category in categories),
    }
    return render(request, 'inventory/category_list.html', {
        'categories': categories, 'query': query, 'category_stats': category_stats,
    })


@login_required
def category_form(request, pk=None):
    organization = _scope(request, PermissionCode.CATEGORY_MANAGE)
    category = get_object_or_404(ProductCategory.objects.filter(tenant=organization), pk=pk) if pk else None
    old = {
        'name': category.name, 'description': category.description, 'measure_unit': category.measure_unit,
        'default_unit_id': category.default_unit_id,
        'allowed_unit_ids': list(category.allowed_units.order_by('pk').values_list('pk', flat=True)),
        'icon': category.icon, 'is_active': category.is_active,
    } if category else {}
    form = ProductCategoryForm(request.POST or None, instance=category, organization=organization)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            form.instance.organization = form.instance.tenant = organization
            obj = form.save()
            audit(organization=organization, actor=request.user, action='inventory.category.saved', obj=obj, old_value=old, new_value={
                'name': obj.name, 'description': obj.description, 'measure_unit': obj.measure_unit,
                'default_unit_id': obj.default_unit_id,
                'allowed_unit_ids': list(obj.allowed_units.order_by('pk').values_list('pk', flat=True)),
                'icon': obj.icon, 'is_active': obj.is_active,
            })
        messages.success(request, 'Category saved.')
        return redirect('inventory:category_list')
    return render(request, 'inventory/category_form.html', {
        'form': form, 'category': category, 'title': 'Edit category' if category else 'Add category',
        'cancel_url': 'inventory:category_list', 'submit_label': 'Save category',
    })


@login_required
@require_POST
def purchase_category_quick_create(request):
    organization = _scope(request, PermissionCode.CATEGORY_MANAGE)
    form = QuickProductCategoryForm(request.POST, organization=organization)
    if not form.is_valid():
        return JsonResponse({'errors': form.errors.get_json_data()}, status=400)

    try:
        with transaction.atomic():
            category = form.save(commit=False)
            category.organization = category.tenant = organization
            category.is_active = True
            category.save()
            category.allowed_units.set([form.cleaned_data['default_unit']])
            audit(
                organization=organization,
                actor=request.user,
                action='inventory.category.created_quick',
                obj=category,
                old_value={},
                new_value={
                    'name': category.name,
                    'default_unit_id': category.default_unit_id,
                    'allowed_unit_ids': [category.default_unit_id],
                    'is_active': category.is_active,
                },
            )
    except IntegrityError:
        if ProductCategory.objects.unscoped().filter(
            tenant=organization,
            name__iexact=form.cleaned_data['name'],
        ).exists():
            errors = {'name': [{'message': 'A category with this name already exists.', 'code': 'unique'}]}
        else:
            errors = {'__all__': [{'message': 'Category could not be saved. Try again.', 'code': 'integrity'}]}
        return JsonResponse({'errors': errors}, status=409)

    return JsonResponse({
        'category': {
            'id': category.pk,
            'text': category.name,
            'default_unit': {
                'id': category.default_unit_id,
                'text': category.default_unit.label,
            },
        },
    }, status=201)


@login_required
def supplier_list(request):
    organization = _scope(request, PermissionCode.SUPPLIER_MANAGE)
    query = request.GET.get('q', '').strip()
    suppliers = Supplier.objects.filter(tenant=organization).annotate(
        confirmed_purchase_count=Count('purchases', filter=Q(purchases__status=Purchase.Status.CONFIRMED)),
        last_purchase_date=Max('purchases__purchase_date', filter=Q(purchases__status=Purchase.Status.CONFIRMED)),
    )
    if query:
        suppliers = suppliers.filter(
            Q(company_name__icontains=query)
            | Q(contact_person__icontains=query)
            | Q(phone__icontains=query)
            | Q(email__icontains=query)
            | Q(tin_vrn__icontains=query)
        )
    return render(request, 'inventory/supplier_list.html', {'suppliers': suppliers, 'query': query})


@login_required
def supplier_form(request, pk=None):
    organization = _scope(request, PermissionCode.SUPPLIER_MANAGE)
    embedded_create = pk is None and request.GET.get('purchase_modal') == '1'
    supplier = get_object_or_404(Supplier.objects.filter(tenant=organization), pk=pk) if pk else None
    old = {'company_name': supplier.company_name, 'is_active': supplier.is_active} if supplier else {}
    form = SupplierForm(request.POST or None, instance=supplier, organization=organization)
    if embedded_create:
        form.fields['is_active'].initial = True
        form.fields['is_active'].disabled = True
    if request.method == 'POST' and form.is_valid():
        obj = form.save(commit=False)
        obj.organization = obj.tenant = organization
        if obj.created_by_id is None:
            obj.created_by = request.user
        obj.save()
        audit(organization=organization, actor=request.user, action='inventory.supplier.saved', obj=obj, old_value=old, new_value={'company_name': obj.company_name, 'is_active': obj.is_active})
        if embedded_create:
            response = render(request, 'inventory/includes/purchase_modal_created.html', {
                'created_record': {'kind': 'supplier', 'id': obj.pk, 'text': str(obj)},
            })
            response['X-Frame-Options'] = 'SAMEORIGIN'
            return response
        messages.success(request, 'Supplier saved.')
        return redirect('inventory:supplier_list')
    response = render(request, 'inventory/supplier_form.html', {
        'form': form, 'supplier': supplier, 'title': 'Edit supplier' if supplier else 'New supplier',
        'embedded_create': embedded_create,
    })
    if embedded_create:
        response['X-Frame-Options'] = 'SAMEORIGIN'
    return response


@login_required
def supplier_detail(request, pk):
    organization = _scope(request, PermissionCode.SUPPLIER_MANAGE)
    supplier = get_object_or_404(Supplier.objects.filter(tenant=organization), pk=pk)
    return render(request, 'inventory/supplier_detail.html', {'supplier': supplier})


@login_required
def supplier_payment_create(request, pk):
    organization = _scope(request, PermissionCode.SUPPLIER_MANAGE)
    supplier = get_object_or_404(Supplier.objects.filter(tenant=organization), pk=pk)
    form = SupplierPaymentForm(request.POST or None, organization=organization)
    if request.method == 'POST' and form.is_valid():
        payment = form.save(commit=False)
        payment.organization = payment.tenant = organization
        payment.supplier = supplier
        payment.created_by = request.user
        payment.save()
        audit(organization=organization, actor=request.user, action='inventory.supplier_payment.recorded', obj=payment, new_value={'supplier_id': supplier.pk, 'amount': str(payment.amount)})
        messages.success(request, 'Supplier payment record saved. This does not post to a general ledger.')
        return redirect('inventory:supplier_detail', pk=supplier.pk)
    return render(request, 'inventory/form.html', {
        'form': form, 'title': f'Record payment — {supplier.company_name}',
        'subtitle': 'This is an operational supplier record and does not post to a general ledger.',
        'supplier': supplier, 'submit_label': 'Record payment',
    })


@login_required
def purchase_list(request):
    organization = _scope(request, PermissionCode.PURCHASE_VIEW)
    purchases = Purchase.objects.filter(tenant=organization).select_related(
        'supplier', 'created_by', 'confirmed_by'
    ).annotate(item_count=Count('lines'))
    status = request.GET.get('status')
    if status in Purchase.Status.values:
        purchases = purchases.filter(status=status)
    query = request.GET.get('q', '').strip()
    if query:
        purchases = purchases.filter(
            Q(reference_number__icontains=query)
            | Q(supplier__company_name__icontains=query)
            | Q(created_by__username__icontains=query)
        )
    return render(request, 'inventory/purchase_list.html', {
        'purchases': purchases, 'statuses': Purchase.Status.choices, 'status': status or '', 'query': query,
    })


PURCHASE_WORKSPACE_ACTIONS = frozenset({'save_continue', 'save_review', 'confirm_receive'})
PURCHASE_PRODUCT_SEARCH_PAGE_SIZE = 20
PURCHASE_PASTE_MAX_ROWS = 200


@login_required
@require_GET
def purchase_rows_template(request):
    _scope(request, PermissionCode.PURCHASES_CREATE)
    output = purchase_rows_template_bytes()
    response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="purchase_rows_template.xlsx"'
    return response


def purchase_rows_template_bytes():
    workbook = Workbook()
    rows = workbook.active
    rows.title = 'Purchase rows'
    rows.append(MIXED_HEADERS)
    rows.freeze_panes = 'A2'
    # Compatibility metadata remains in the workbook, outside the ordinary UI.
    rows.column_dimensions['O'].hidden = True
    workbook.properties.version = '2'
    widths = (24, 32, 16, 18, 28, 18, 18, 24, 24, 16, 16, 24, 18, 42, 24)
    for index, width in enumerate(widths, start=1):
        letter = rows.cell(1, index).column_letter
        rows.column_dimensions[letter].width = width
        cell = rows.cell(1, index)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4F35D7' if index <= 3 else '64748B')
        cell.alignment = Alignment(wrap_text=True)
    for row in rows.iter_rows(min_row=2, max_row=201, min_col=1, max_col=len(MIXED_HEADERS)):
        for column in (0, 13, 14):
            row[column].number_format = '@'
        for column in (2, 3, 6):
            row[column].number_format = '0.######'
        row[12].number_format = 'yyyy-mm-dd'
    instructions = workbook.create_sheet('Instructions')
    instructions.column_dimensions['A'].width = 105
    lines = [
        'Purchase rows template — instructions',
        'Purchase rows sheet has one header and blank data rows. Instructions/examples on this sheet are never imported.',
        'Existing example: enter an active stock SKU; other product attributes are taken from the catalog.',
        'New example: leave SKU blank, enter product name, category, unit, selling price and tracking choices.',
        'Unknown supplied SKU requires explicit Create new or an existing-product selection.',
        'Choose Use existing or Create new explicitly for items without a known SKU.',
        'Quantity and unit cost use that product’s selected sales/stock unit.',
        'Use numbers such as 1250.50 or 1,250.50 without currency symbols. Quantity and unit cost support six decimals; selling price supports two.',
        'Format SKU and serial numbers as text to preserve leading zeros.',
        'Use YYYY-MM-DD for expiry. Enter comma-separated serial numbers in one cell.',
        'Serialized products need one unique serial for each whole unit. Other tracking fields follow purchase validation.',
        f'Paste at most {PURCHASE_PASTE_MAX_ROWS} rows. Include headers for mixed products; six-column legacy rows may omit them.',
        'Previewing does not save a purchase or change stock.',
        '',
        'ILLUSTRATIVE EXAMPLES — these are not purchase rows',
        'Existing: EXAMPLE-SKU | Router | 2 | 12.345678',
        'New: blank SKU | New router | 2 | 12.345678 | Devices | pc | 90.00',
    ]
    for line in lines:
        instructions.append([line])
    instructions['A1'].font = Font(bold=True, size=14)
    instructions['A15'].font = Font(bold=True, color='9A3412')
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@login_required
@require_GET
@never_cache
def purchase_product_search(request):
    """Return a small tenant-scoped product page for the purchase combobox."""
    organization = _scope(request, PermissionCode.PURCHASES_CREATE)
    query = request.GET.get('q', '').strip()[:100]
    try:
        page = min(max(int(request.GET.get('page', '1')), 1), 1000)
    except (TypeError, ValueError):
        page = 1

    products = Product.objects.filter(
        tenant=organization,
        is_active=True,
        item_type=Product.ItemType.PHYSICAL,
        track_stock=True,
    )
    if query:
        products = products.filter(
            Q(name__icontains=query)
            | Q(sku__icontains=query)
            | Q(brand__icontains=query)
            | Q(model_number__icontains=query)
        )
    products = products.select_related('sales_unit').order_by('name', 'sku', 'pk')
    start = (page - 1) * PURCHASE_PRODUCT_SEARCH_PAGE_SIZE
    rows = list(products[start:start + PURCHASE_PRODUCT_SEARCH_PAGE_SIZE + 1])
    more = len(rows) > PURCHASE_PRODUCT_SEARCH_PAGE_SIZE

    can_view_costs = has_tenant_permission(
        request.user, organization, PermissionCode.COST_REPORT_VIEW, membership=request.membership,
    )
    return JsonResponse({
        'results': [
            {
                'id': product.pk,
                'text': str(product),
                'name': product.name,
                'sku': product.sku,
                'unit': product.sales_unit.label if product.sales_unit_id else product.get_measure_unit_display(),
                'unit_id': product.sales_unit_id,
                'brand': product.brand,
                'model_number': product.model_number,
                'serialized': product.is_serialized,
                'expiry': product.track_expiry,
                **({'buying_price': str(product.buying_price)} if can_view_costs else {}),
            }
            for product in rows[:PURCHASE_PRODUCT_SEARCH_PAGE_SIZE]
        ],
        'pagination': {'page': page, 'more': more},
    })


@login_required
@require_POST
@never_cache
def serial_check(request):
    """Bounded, non-mutating feedback; never reserves a serial."""
    mode = request.POST.get('mode', '')
    if mode not in ('purchase', 'adjustment'):
        return JsonResponse({'error': 'Invalid serial-check context.'}, status=400)
    organization = _scope(request, PermissionCode.PURCHASES_CREATE if mode == 'purchase' else PermissionCode.STOCK_ADJUST)
    raw = request.POST.get('serials', '')
    if len(raw) > 100000:
        return JsonResponse({'error': 'Check at most 1,000 serials at a time.'}, status=400)
    try:
        serials = json.loads(raw)
    except (TypeError, ValueError):
        serials = None
    if (not isinstance(serials, list) or len(serials) > 1000
            or any(not isinstance(value, str) or not value.strip() or '\x00' in value or len(value.strip()) > 160 for value in serials)):
        return JsonResponse({'error': 'Provide at most 1,000 serials, each up to 160 characters.'}, status=400)
    normalized = {value.strip().upper() for value in serials}
    if any(len(value) > 160 for value in normalized):
        return JsonResponse({'error': 'Normalized serials must be at most 160 characters.'}, status=400)
    from .serial_validation import existing_serials
    if mode == 'adjustment':
        direction = request.POST.get('direction')
        if direction not in ('increase', 'decrease'):
            return JsonResponse({'error': 'Choose the stock adjustment direction.'}, status=400)
        try:
            product_id = int(request.POST.get('product', ''))
        except (TypeError, ValueError):
            product_id = None
        product = get_object_or_404(Product.objects.filter(
            tenant=organization, is_active=True, item_type=Product.ItemType.PHYSICAL, track_stock=True,
        ), pk=product_id)
        if direction == 'decrease':
            available = set(StockUnit.objects.filter(
                tenant=organization, product=product, status=StockUnit.Status.AVAILABLE,
                serial_number__in=normalized,
            ).values_list('serial_number', flat=True))
            return JsonResponse({'conflicts': sorted(normalized - available), 'reason': 'unavailable'})
    return JsonResponse({'conflicts': sorted(existing_serials(organization=organization, serials=normalized)), 'reason': 'exists'})


@login_required
@require_POST
def purchase_rows_preview(request):
    """Shared XLSX/TSV preview. No products, purchases or stock are written."""
    organization = _scope(request, PermissionCode.PURCHASES_CREATE)
    try:
        operation = request.POST.get('operation', 'preview')
        if operation not in ('preview', 'inspect_columns'):
            raise ImportInputError('Choose a supported purchase import action.')
        header_choice = request.POST.get('has_header')
        if header_choice is not None and header_choice not in ('true', 'false'):
            raise ImportInputError('Confirm whether the first row contains column headings.')
        has_header = None if header_choice is None else header_choice == 'true'
        if operation == 'inspect_columns':
            if 'review_rows' in request.POST:
                raise ImportInputError('Choose columns from the original spreadsheet input.')
            return JsonResponse({'column_mapping': inspect_input(
                text=request.POST.get('rows', ''), upload=request.FILES.get('file'), has_header=has_header,
            )})
        column_mapping = None
        if 'column_mapping' in request.POST:
            encoded_mapping = request.POST['column_mapping']
            if len(encoded_mapping) > 2000:
                raise ImportInputError('Column mapping is too large. Choose columns again.')
            try:
                column_mapping = json.loads(encoded_mapping)
            except (ValueError, TypeError) as exc:
                raise ImportInputError('Column mapping is invalid. Choose columns again.') from exc
            if not isinstance(column_mapping, list):
                raise ImportInputError('Column mapping must be a list of field selections.')
        rows = (submitted_rows(request.POST['review_rows']) if 'review_rows' in request.POST
                else parse_input(text=request.POST.get('rows', ''), upload=request.FILES.get('file'),
                                 input_format=request.POST.get('format', 'auto'),
                                 column_mapping=column_mapping, has_header=has_header))
        results = validate_rows(rows, organization=organization)
    except ImportColumnMappingRequired as exc:
        return JsonResponse({'error': str(exc), 'code': 'column_mapping_required',
                             'column_mapping': exc.structure}, status=400)
    except ImportInputError as exc:
        return JsonResponse({'error': str(exc)}, status=400)
    counts = {status: sum(row['status'] == status for row in results)
              for status in ('matched', 'new', 'unresolved', 'invalid', 'excluded')}
    detailed_summary = {
        'total': len(results), 'valid': sum(row['valid'] for row in results),
        'matched': counts['matched'], 'new': counts['new'],
        'needs_attention': counts['unresolved'] + counts['invalid'], 'excluded': counts['excluded'],
    }
    legacy = 'review_rows' not in request.POST and not request.FILES and not any(
        'product_name' in row for row in rows
    )
    return JsonResponse({'rows': results, 'summary': (
        {'total': detailed_summary['total'], 'valid': detailed_summary['valid']} if legacy else detailed_summary
    ), 'review_summary': detailed_summary,
        'format_label': ('Custom column mapping' if column_mapping is not None else
                         'Mixed product columns' if any('product_name' in row for row in rows) else 'Legacy SKU columns'),
        'ignored_columns': column_mapping.count('__ignore__') if column_mapping is not None else 0})


@login_required
@require_POST
def purchase_product_quick_create(request):
    from audit.models import AuditLog
    organization = _scope(request, PermissionCode.PRODUCTS_CREATE)
    form = QuickPurchaseProductForm(request.POST, organization=organization)
    key = request.headers.get('Idempotency-Key', '')
    if len(key) > 128:
        return JsonResponse({'message': 'Idempotency-Key must be at most 128 characters.'}, status=400)
    from products.workspace import submission_fingerprint
    fingerprint = submission_fingerprint({k: v for k, v in request.POST.lists() if k != 'csrfmiddlewaretoken'}, {})
    try:
        with transaction.atomic():
            type(organization).objects.select_for_update().get(pk=organization.pk)
            if key:
                prior = AuditLog.objects.unscoped().filter(tenant=organization, actor=request.user,
                    action='inventory.product.created_quick', metadata__key=key).first()
                if prior:
                    if prior.metadata.get('fingerprint') != fingerprint:
                        return JsonResponse({'message': 'Idempotency key was already used with different content.'}, status=409)
                    product = get_object_or_404(Product.objects.unscoped(), tenant=organization, pk=prior.object_id)
                    return _quick_product_response(product, status=200)
            if not form.is_valid():
                return JsonResponse({'errors': form.errors.get_json_data()}, status=400)
            product = form.save(commit=True)
            audit(
                organization=organization, actor=request.user,
                action='inventory.product.created_quick', obj=product,
                metadata={'key': key, 'fingerprint': fingerprint} if key else {},
                new_value={
                    'name': product.name, 'sku': product.sku,
                    'catalog_category_id': product.catalog_category_id,
                    'sales_unit_id': product.sales_unit_id,
                    'is_serialized': product.is_serialized,
                    'track_expiry': product.track_expiry,
                },
            )
    except (IntegrityError, ValidationError):
        return JsonResponse({'errors': {'__all__': [{
            'message': 'Product could not be saved safely. Refresh and try again.', 'code': 'integrity',
        }]}}, status=409)
    return _quick_product_response(product, status=201)


def _quick_product_response(product, *, status):
    return JsonResponse({'product': {
        'id': product.pk, 'text': str(product), 'name': product.name, 'sku': product.sku,
        'unit': product.sales_unit.label, 'serialized': product.is_serialized,
        'expiry': product.track_expiry,
    }}, status=status)


@login_required
@require_POST
def purchase_supplier_quick_create(request):
    organization = _scope(request, PermissionCode.SUPPLIER_MANAGE)
    form = QuickSupplierForm(request.POST, organization=organization)
    if not form.is_valid():
        return JsonResponse({'errors': form.errors.get_json_data()}, status=400)
    try:
        with transaction.atomic():
            supplier = form.save(commit=False)
            supplier.organization = supplier.tenant = organization
            supplier.created_by = request.user
            supplier.is_active = True
            supplier.save()
            audit(
                organization=organization,
                actor=request.user,
                action='inventory.supplier.created_quick',
                obj=supplier,
                new_value={'company_name': supplier.company_name, 'phone': supplier.phone, 'is_active': True},
            )
    except IntegrityError:
        duplicate = Supplier.objects.unscoped().filter(
            tenant=organization,
            company_name__iexact=form.cleaned_data['company_name'],
        ).exists()
        return JsonResponse({
            'errors': (
                {'company_name': [{
                    'message': 'This supplier already exists in your organization.',
                    'code': 'unique',
                }]}
                if duplicate else
                {'__all__': [{'message': 'Supplier could not be saved safely. Try again.', 'code': 'integrity'}]}
            ),
        }, status=409)
    return JsonResponse({
        'supplier': {'id': supplier.pk, 'text': supplier.company_name, 'phone': supplier.phone},
    }, status=201)


def _purchase_product_meta(organization, formset):
    product_ids = {
        form.instance.product_id
        for form in formset.forms
        if form.instance and form.instance.product_id
    }
    return {
        str(product.pk): {
            'name': product.name,
            'serialized': product.is_serialized,
            'expiry': product.track_expiry,
            'sku': product.sku,
            'base_unit': product.sales_unit.label if product.sales_unit_id else product.get_measure_unit_display(),
        }
        for product in Product.objects.filter(
            tenant=organization,
            pk__in=product_ids,
            is_active=True,
            item_type=Product.ItemType.PHYSICAL,
            track_stock=True,
        ).select_related('sales_unit').order_by('name')
    }


def _purchase_workspace_summary(purchase, formset):
    if formset.is_bound:
        valid_lines = [
            line_form.instance
            for line_form in formset.forms
            if line_form.cleaned_data
            and not line_form.errors
            and not line_form.cleaned_data.get('DELETE')
            and line_form.cleaned_data.get('product')
        ]
    elif purchase.pk:
        valid_lines = list(purchase.lines.select_related('product'))
        from .purchase_batch import proposal_lines
        valid_lines.extend(proposal_lines(purchase))
    else:
        valid_lines = []
    return len(valid_lines), sum((line.line_total for line in valid_lines), Decimal('0.00'))


def _render_purchase_workspace(request, *, organization, purchase, form, formset):
    from .purchase_batch import proposal_rows
    staged_rows = proposal_rows(purchase) if request.method != 'POST' else []
    line_count, authoritative_total = _purchase_workspace_summary(purchase, formset)
    return render(request, 'inventory/purchase_form.html', {
        'form': form,
        'formset': formset,
        'purchase': purchase,
        'workspace_title': 'Edit Purchase' if purchase.pk else 'New Purchase',
        'product_meta': _purchase_product_meta(organization, formset),
        'product_search_url': reverse('inventory:purchase_product_search'),
        'purchase_rows_preview_url': reverse('inventory:purchase_rows_preview'),
        'import_token': request.POST.get('import_token') or str(uuid.uuid4()),
        'import_rows': request.POST.get('import_rows', '') if request.method == 'POST' else (json.dumps(staged_rows) if staged_rows else ''),
        'supplier_quick_create_url': reverse('inventory:purchase_supplier_quick_create'),
        'category_quick_create_url': reverse('inventory:purchase_category_quick_create'),
        'product_quick_create_url': reverse('inventory:purchase_product_quick_create'),
        'quick_product_categories': ProductCategory.objects.filter(tenant=organization, is_active=True).prefetch_related('allowed_units'),
        'quick_product_units': UnitOfMeasure.objects.filter(tenant=organization, is_active=True).order_by('name'),
        'can_create_import_product': has_tenant_permission(
            request.user, organization, PermissionCode.PRODUCTS_CREATE, membership=request.membership
        ),
        'valid_line_count': line_count,
        'authoritative_total': authoritative_total.quantize(Decimal('0.01')),
        'can_confirm_purchase': has_tenant_permission(
            request.user, organization, PermissionCode.PURCHASE_CONFIRM, membership=request.membership
        ),
    })


def _save_purchase_draft(*, form, formset, organization, actor, creating,
                         import_rows=None, import_token=None, import_fingerprint=None, membership=None):
    with transaction.atomic():
        if import_rows:
            type(organization).objects.select_for_update().get(pk=organization.pk)
            previous = PurchaseImportSubmission.objects.filter(tenant=organization, token=import_token).first()
            if previous:
                if previous.fingerprint != import_fingerprint:
                    raise ImportInputError('This import token was already used with different content. Review and apply a new batch.')
                return previous.purchase
        purchase = form.save(commit=False)
        purchase.organization = purchase.tenant = organization
        if creating:
            purchase.created_by = actor
            if form.cleaned_data['auto_generated_reference'] == purchase.reference_number:
                purchase.reference_number = PurchaseReferenceNumberService.next_number(
                    organization=organization,
                    purchase_date=purchase.purchase_date,
                )
        purchase.save()
        formset.instance = purchase
        formset.save()
        # The workspace submits the complete proposal snapshot on edit.
        purchase.product_proposals.filter(tenant=organization).delete()
        if import_rows:
            apply_purchase_batch(purchase=purchase, rows=import_rows, organization=organization,
                                 actor=actor, membership=membership)
        from .purchase_batch import proposal_lines
        total = sum((line.line_total for line in list(purchase.lines.all()) + proposal_lines(purchase)), Decimal('0.00'))
        total = total.quantize(Decimal('0.01'))
        DecimalValidator(14, 2)(total)
        Purchase.objects.filter(tenant=organization, pk=purchase.pk).update(total_cost=total)
        purchase.total_cost = total
        audit(
            organization=organization,
            actor=actor,
            action='inventory.purchase.created' if creating else 'inventory.purchase.updated',
            obj=purchase,
            new_value={'reference_number': purchase.reference_number, 'total_cost': str(total)},
        )
        if import_rows:
            PurchaseImportSubmission.objects.create(
                organization=organization, tenant=organization, token=import_token,
                fingerprint=import_fingerprint, purchase=purchase,
            )
    return purchase


def _import_request(request, organization):
    raw = request.POST.get('import_rows', '')
    if not raw:
        return [], None, None
    rows = submitted_rows(raw)
    if len(rows) > PURCHASE_PASTE_MAX_ROWS:
        raise ImportInputError('The batch exceeds 200 rows.')
    try:
        manual_count = int(request.POST.get('lines-TOTAL_FORMS', '0'))
    except ValueError as exc:
        raise ImportInputError('Purchase row count is invalid.') from exc
    if manual_count + len(rows) > 250:
        raise ImportInputError('The workspace supports at most 250 total rows.')
    try:
        token = uuid.UUID(request.POST.get('import_token', ''))
    except (ValueError, AttributeError) as exc:
        raise ImportInputError('Import token is missing or invalid. Review the batch again.') from exc
    pairs = sorted((key, request.POST.getlist(key)) for key in request.POST if key != 'csrfmiddlewaretoken')
    fingerprint = hashlib.sha256(json.dumps(pairs, sort_keys=True).encode('utf-8')).hexdigest()
    return rows, token, fingerprint


def _import_replay(organization, token, fingerprint, action, expected_purchase_id=None):
    if not token:
        return None
    receipt = PurchaseImportSubmission.objects.filter(tenant=organization, token=token).select_related('purchase').first()
    if not receipt:
        return None
    if receipt.fingerprint != fingerprint:
        raise ImportInputError('This import token was already saved with different content. Review a new batch.')
    if expected_purchase_id and receipt.purchase_id != expected_purchase_id:
        raise ImportInputError('This import token belongs to another purchase. Review a new batch.')
    return _purchase_action_redirect(receipt.purchase, action)


def _purchase_action_redirect(purchase, action):
    if action == 'save_continue':
        return redirect('inventory:purchase_edit', pk=purchase.pk)
    if action == 'confirm_receive':
        return redirect(f"{reverse('inventory:purchase_detail', kwargs={'pk': purchase.pk})}?receive=1")
    return redirect('inventory:purchase_detail', pk=purchase.pk)


@login_required
@transaction.atomic
def purchase_create(request):
    organization = _scope(request, PermissionCode.PURCHASES_CREATE)
    import_rows, import_token, import_fingerprint, import_error = [], None, None, None
    if request.method == 'POST':
        try:
            import_rows, import_token, import_fingerprint = _import_request(request, organization)
            if import_rows:
                type(organization).objects.select_for_update().get(pk=organization.pk)
                replay = _import_replay(organization, import_token, import_fingerprint, request.POST.get('action', ''))
                if replay:
                    return replay
        except ImportInputError as exc:
            import_error = str(exc)
    purchase = Purchase(organization=organization, tenant=organization, created_by=request.user)
    form_initial = None
    # Previewing does not consume a sequence number. The hidden marker lets a
    # successful POST atomically replace an untouched preview with the actual
    # next number, while preserving a user-entered supplier reference.
    if request.method == 'GET':
        purchase.reference_number = PurchaseReferenceNumberService.preview_next_number(organization=organization)
        form_initial = {'auto_generated_reference': purchase.reference_number}
    form = PurchaseForm(request.POST or None, instance=purchase, organization=organization, initial=form_initial)
    formset = PurchaseLinesFormSet(request.POST or None, instance=purchase, organization=organization,
                                   import_present=bool(import_rows))
    if request.method == 'POST':
        if import_error:
            form.add_error(None, import_error)
        action = request.POST.get('action', '')
        form_valid = form.is_valid()
        formset_valid = formset.is_valid()
        if action not in PURCHASE_WORKSPACE_ACTIONS:
            form.add_error(None, 'Choose a valid purchase action and try again.')
        elif form_valid and formset_valid and not import_error:
            try:
                purchase = _save_purchase_draft(
                    form=form, formset=formset, organization=organization, actor=request.user, creating=True,
                    import_rows=import_rows, import_token=import_token, import_fingerprint=import_fingerprint,
                    membership=request.membership,
                )
            except IntegrityError:
                form.add_error('reference_number', 'This purchase reference already exists. Choose another reference.')
            except (ImportInputError, ValidationError) as exc:
                form.add_error(None, str(exc))
            else:
                if action == 'save_continue':
                    messages.success(request, 'Purchase draft saved. Continue editing when ready.')
                elif action == 'confirm_receive':
                    messages.success(request, 'Purchase draft saved. Review every line before confirming receipt.')
                else:
                    messages.success(request, 'Purchase draft saved for review. Stock has not changed.')
                return _purchase_action_redirect(purchase, action)
    return _render_purchase_workspace(
        request, organization=organization, purchase=purchase, form=form, formset=formset
    )


@login_required
@transaction.atomic
def purchase_edit(request, pk):
    organization = _scope(request, PermissionCode.PURCHASES_CREATE)
    purchase = get_object_or_404(Purchase.objects.filter(tenant=organization), pk=pk)
    if not has_tenant_permission(request.user, organization, PermissionCode.PURCHASE_VIEW,
                                 membership=request.membership) and purchase.created_by_id != request.user.pk:
        raise Http404
    if purchase.status != Purchase.Status.DRAFT:
        messages.error(request, 'Only draft purchases can be edited.')
        return redirect('inventory:purchase_detail', pk=purchase.pk)
    import_rows, import_token, import_fingerprint, import_error = [], None, None, None
    if request.method == 'POST':
        try:
            import_rows, import_token, import_fingerprint = _import_request(request, organization)
            if import_rows:
                type(organization).objects.select_for_update().get(pk=organization.pk)
                replay = _import_replay(organization, import_token, import_fingerprint,
                                        request.POST.get('action', ''), expected_purchase_id=purchase.pk)
                if replay:
                    return replay
        except ImportInputError as exc:
            import_error = str(exc)
    form = PurchaseForm(request.POST or None, instance=purchase, organization=organization)
    formset = PurchaseLinesFormSet(request.POST or None, instance=purchase, organization=organization,
                                   import_present=bool(import_rows))
    if request.method == 'POST':
        if import_error:
            form.add_error(None, import_error)
        action = request.POST.get('action', '')
        form_valid = form.is_valid()
        formset_valid = formset.is_valid()
        if action not in PURCHASE_WORKSPACE_ACTIONS:
            form.add_error(None, 'Choose a valid purchase action and try again.')
        elif form_valid and formset_valid and not import_error:
            try:
                purchase = _save_purchase_draft(
                    form=form, formset=formset, organization=organization, actor=request.user, creating=False,
                    import_rows=import_rows, import_token=import_token, import_fingerprint=import_fingerprint,
                    membership=request.membership,
                )
            except IntegrityError:
                form.add_error('reference_number', 'This purchase reference already exists. Choose another reference.')
            except (ImportInputError, ValidationError) as exc:
                form.add_error(None, str(exc))
            else:
                if action == 'save_continue':
                    messages.success(request, 'Purchase draft updated. Continue editing when ready.')
                elif action == 'confirm_receive':
                    messages.success(request, 'Purchase draft updated. Review every line before confirming receipt.')
                else:
                    messages.success(request, 'Purchase draft updated for review. Stock has not changed.')
                return _purchase_action_redirect(purchase, action)
    return _render_purchase_workspace(
        request, organization=organization, purchase=purchase, form=form, formset=formset
    )


@login_required
def purchase_cancel(request, pk):
    organization = _scope(request, PermissionCode.PURCHASE_CONFIRM)
    if request.method != 'POST':
        raise Http404
    try:
        purchase = InventoryService.cancel_purchase(organization=organization, purchase_id=pk, actor=request.user)
    except InventoryError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, f'Purchase {purchase.reference_number} cancelled. Stock was not changed.')
    return redirect('inventory:purchase_detail', pk=pk)


@login_required
@ensure_csrf_cookie
@never_cache
def purchase_detail(request, pk):
    try:
        organization = _scope(request, PermissionCode.PURCHASE_VIEW)
        can_view_all = True
    except PermissionDenied:
        organization = _scope(request, PermissionCode.PURCHASES_CREATE)
        can_view_all = False
    purchase = get_object_or_404(Purchase.objects.filter(tenant=organization).select_related('supplier', 'created_by', 'confirmed_by'), pk=pk)
    if not can_view_all and purchase.created_by_id != request.user.pk:
        raise Http404
    # Lines are the source of truth for a draft receipt. This also renders a
    # correct total for historical drafts that predate draft-total syncing.
    purchase_lines = list(purchase.lines.select_related('product'))
    from .purchase_batch import proposal_lines
    purchase_lines.extend(proposal_lines(purchase))
    line_items_total = sum((line.line_total for line in purchase_lines), Decimal('0.00'))
    incoming_by_product = {}
    for line in purchase_lines:
        totals = incoming_by_product.setdefault(line.product_id, {'quantity': Decimal('0'), 'cost': Decimal('0')})
        totals['quantity'] += line.quantity
        totals['cost'] += line.line_total
    balances = {
        balance.product_id: balance
        for balance in InventoryBalance.objects.filter(
            tenant=organization, product_id__in=incoming_by_product,
        )
    }
    pricing_reviews = []
    seen_products = set()
    for line in purchase_lines:
        if line.product_id in seen_products:
            continue
        seen_products.add(line.product_id)
        product = line.product
        if purchase.status == Purchase.Status.CONFIRMED:
            projected_floor = cost_floor_for(product)
            warnings = product.pricing_warnings
        else:
            incoming = incoming_by_product[line.product_id]
            balance = balances.get(line.product_id)
            current_quantity = balance.quantity if balance and balance.quantity > 0 else Decimal('0')
            current_cost = balance.average_cost if current_quantity > 0 else Decimal(product.buying_price or 0)
            total_quantity = current_quantity + incoming['quantity']
            projected_floor = (
                ((current_quantity * current_cost) + incoming['cost']) / total_quantity
                if total_quantity > 0 else current_cost
            ).quantize(Decimal('0.000001'))
            prices = [('Standard', product.selling_price), ('Technician', product.effective_technician_price)]
            if product.allow_wholesale and product.wholesale_price is not None:
                prices.append(('Wholesale', product.wholesale_price))
            warnings = [label for label, price in prices if price is not None and Decimal(price) <= projected_floor]
        if warnings:
            pricing_reviews.append({
                'product': product,
                'projected_floor': projected_floor,
                'categories': warnings,
            })
    return render(request, 'inventory/purchase_detail.html', {
        'purchase': purchase,
        'can_edit_purchase': purchase.status == Purchase.Status.DRAFT and has_tenant_permission(
            request.user, organization, PermissionCode.PURCHASES_CREATE, membership=request.membership
        ) and (can_view_all or purchase.created_by_id == request.user.pk),
        'purchase_lines': purchase_lines,
        'line_items_total': line_items_total,
        'product_type_count': len(incoming_by_product),
        'serialized_line_count': sum(1 for line in purchase_lines if line.product.is_serialized),
        'expiry_line_count': sum(1 for line in purchase_lines if line.product.track_expiry),
        'pricing_reviews': pricing_reviews,
    })


@login_required
def purchase_confirm(request, pk):
    organization = _scope(request, PermissionCode.PURCHASE_CONFIRM)
    if request.method != 'POST':
        raise Http404
    try:
        purchase = InventoryService.confirm_purchase(organization=organization, purchase_id=pk, actor=request.user)
        messages.success(request, f'Purchase {purchase.reference_number} confirmed and stock received.')
        if has_tenant_permission(request.user, organization, PermissionCode.COST_REPORT_VIEW, membership=request.membership):
            warnings = {
                line.product.name: line.product.pricing_warnings
                for line in purchase.lines.select_related('product')
                if line.product.pricing_warnings
            }
            if warnings:
                labels = '; '.join(f"{name}: {', '.join(categories)}" for name, categories in warnings.items())
                messages.warning(request, f'Pricing review required after receiving stock — {labels}. Affected sales are blocked.')
    except (InventoryError, BillingServiceError) as exc:
        messages.error(request, str(exc))
    return redirect('inventory:purchase_detail', pk=pk)


@login_required
def stock_list(request):
    organization = _scope(request, PermissionCode.STOCK_VIEW)
    last_movement = StockMovement.objects.filter(
        tenant=organization, product_id=OuterRef('product_id')
    ).order_by('-created_at').values('created_at')[:1]
    balances = InventoryBalance.objects.filter(tenant=organization).select_related(
        'product', 'product__catalog_category'
    ).annotate(last_movement_at=Subquery(last_movement))
    query = request.GET.get('q', '').strip()
    state = request.GET.get('state', '')
    if query:
        balances = balances.filter(Q(product__name__icontains=query) | Q(product__sku__icontains=query))
    if state == 'low':
        balances = balances.filter(quantity__gt=0, quantity__lte=F('product__reorder_threshold'))
    elif state == 'out':
        balances = balances.filter(quantity=0)
    elif state == 'serialized':
        balances = balances.filter(product__is_serialized=True)
    elif state == 'expiring':
        cutoff = timezone.now().date() + timedelta(days=30)
        balances = balances.filter(
            product__track_expiry=True,
            product__stock_units__status=StockUnit.Status.AVAILABLE,
            product__stock_units__expiry_date__lte=cutoff,
        ).distinct()
    return render(request, 'inventory/stock_list.html', {
        'balances': balances.order_by('product__name'), 'query': query, 'state': state,
    })


@login_required
def movement_list(request):
    organization = _scope(request, PermissionCode.STOCK_MOVEMENT_VIEW)
    movements = StockMovement.objects.filter(tenant=organization).select_related(
        'product', 'created_by', 'purchase_line__purchase', 'billing_line__document', 'adjustment'
    )
    product_id = request.GET.get('product')
    if product_id:
        movements = movements.filter(product_id=product_id)
    query = request.GET.get('q', '').strip()
    movement_type = request.GET.get('type', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    if query:
        movements = movements.filter(Q(product__name__icontains=query) | Q(product__sku__icontains=query))
    if movement_type in StockMovement.MovementType.values:
        movements = movements.filter(movement_type=movement_type)
    if date_from:
        movements = movements.filter(created_at__date__gte=date_from)
    if date_to:
        movements = movements.filter(created_at__date__lte=date_to)
    page_obj = Paginator(movements, 50).get_page(request.GET.get('page'))
    return render(request, 'inventory/movement_list.html', {
        'movements': page_obj.object_list,
        'page_obj': page_obj,
        'products': Product.objects.filter(tenant=organization, track_stock=True).order_by('name'),
        'movement_types': StockMovement.MovementType.choices,
        'product_id': str(product_id or ''),
        'query': query,
        'movement_type': movement_type,
        'date_from': date_from,
        'date_to': date_to,
    })


@login_required
def stock_adjust(request):
    organization = _scope(request, PermissionCode.STOCK_ADJUST)
    initial = {'product': request.GET.get('product')} if request.method == 'GET' and request.GET.get('product') else None
    form = StockAdjustmentForm(request.POST or None, organization=organization, initial=initial)
    if request.method == 'POST' and form.is_valid():
        try:
            InventoryService.adjust_stock(
                organization=organization,
                product_id=form.cleaned_data['product'].pk,
                quantity_delta=form.cleaned_data['quantity_delta'],
                reason=form.cleaned_data['reason'],
                notes=form.cleaned_data['notes'],
                serial_numbers=form.serial_list(),
                actor=request.user,
            )
            messages.success(request, 'Stock adjustment recorded.')
            return redirect('inventory:stock_list')
        except InventoryError as exc:
            form.add_error(None, str(exc))
    products = list(form.fields['product'].queryset.select_related('inventory_balance'))
    stock_levels = {str(product.pk): str(product.available_stock) for product in products}
    return render(request, 'inventory/stock_adjust.html', {
        'form': form, 'title': 'Stock adjustment', 'stock_levels': stock_levels,
    })


@login_required
def cart_list(request):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    base = visible_carts(organization=organization, actor=request.user, membership=request.membership)
    now = timezone.now()
    cutoff = now - timedelta(days=7)
    counts = base.aggregate(
        active=Count('pk', filter=Q(status=Cart.Status.DRAFT)),
        stale=Count('pk', filter=Q(status=Cart.Status.DRAFT, updated_at__lte=cutoff)),
        converted=Count('pk', filter=Q(status=Cart.Status.CONVERTED)),
        abandoned=Count('pk', filter=Q(status=Cart.Status.ABANDONED)),
        all=Count('pk'),
    )
    tab = request.GET.get('tab', request.GET.get('status', 'active')).strip() or 'active'
    if tab == 'draft':
        tab = 'active'
    if tab not in counts:
        tab = 'active'
    tabs = [(value, label, counts[value]) for value, label in (
        ('active', 'Active'), ('stale', 'Stale drafts'), ('converted', 'Converted'),
        ('abandoned', 'Abandoned'), ('all', 'All history'),
    )]
    carts = base
    if tab in {'active', 'stale'}:
        carts = carts.filter(status=Cart.Status.DRAFT)
        if tab == 'stale':
            carts = carts.filter(updated_at__lte=cutoff)
    elif tab != 'all':
        carts = carts.filter(status=tab)
    query = request.GET.get('q', '').strip()[:200]
    if query:
        search = (Q(customer__name__icontains=query) | Q(walk_in_name__icontains=query)
                  | Q(quotation__number__icontains=query) | Q(invoice__number__icontains=query))
        number = query.lstrip('#')
        if number.isascii() and number.isdigit() and len(number) <= 18:
            search |= Q(pk=int(number))
        carts = carts.filter(search)
    # Options come only from the same authorized base queryset as rows/counts.
    owners = base.order_by('created_by_id').values('created_by_id', 'created_by__username').distinct()
    owner = request.GET.get('owner', '')
    if owner.isascii() and owner.isdigit() and len(owner) <= 18:
        carts = carts.filter(created_by_id=int(owner))
    dates = {}
    filter_errors = []
    for key in ('from', 'to'):
        raw = request.GET.get(key, '')
        dates[key] = raw
        if raw:
            try:
                parsed = date.fromisoformat(raw)
                from datetime import datetime, time
                boundary = timezone.make_aware(datetime.combine(parsed, time.min))
                if key == 'from':
                    carts = carts.filter(updated_at__gte=boundary)
                else:
                    carts = carts.filter(updated_at__lt=boundary + timedelta(days=1))
            except (ValueError, OverflowError):
                filter_errors.append('Enter a valid last activity date (YYYY-MM-DD).')
                carts = carts.none()
    carts = carts.select_related('customer', 'created_by', 'quotation', 'invoice').annotate(
        item_count=Count('lines'),
        has_business_sale=Exists(InventorySale.objects.unscoped().filter(cart_id=OuterRef('pk'))),
    ).prefetch_related(Prefetch('lines', queryset=CartLine.objects.filter(tenant=organization).select_related('product'), to_attr='list_lines'))
    page = Paginator(carts.order_by('-updated_at', '-pk'), 25).get_page(request.GET.get('page'))
    for cart in page.object_list:
        cart.is_stale = stale(cart, now=now)
        cart.needs_review = cart.status == Cart.Status.DRAFT and (cart.quotation_id or cart.invoice_id or cart.converted_at or cart.has_business_sale)
        try:
            _, _, cart.display_total = BillingService.compute_totals(
                tax_rate=cart.tax_rate, line_items=cart.list_lines, discount_amount=cart.discount_amount,
            )
        except BillingServiceError:
            cart.display_total = None
    params = request.GET.copy()
    params.pop('page', None)
    tab_params = params.copy()
    tab_params.pop('tab', None)
    tab_params.pop('status', None)
    return render(request, 'inventory/cart_list.html', {
        'carts': page.object_list, 'page_obj': page, 'tabs': tabs, 'tab': tab,
        'query': query, 'owner': owner, 'owners': owners, 'dates': dates,
        'filter_errors': filter_errors, 'filter_query': params.urlencode(),
        'tab_query': tab_params.urlencode(),
    })


@login_required
def cart_create(request):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    if request.method != 'POST':
        raise Http404

    cart = CartService.start(organization=organization, actor=request.user, membership=request.membership)
    return redirect('inventory:cart_detail', pk=cart.pk)


@login_required
@ensure_csrf_cookie
@cart_access
def cart_detail(request, pk, *, checkout_forms=None):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    cart = get_object_or_404(Cart.objects.filter(tenant=organization).select_related('customer', 'quotation', 'invoice'), pk=pk)
    cart_form = checkout_forms[0] if checkout_forms else CartForm(request.POST or None, instance=cart, organization=organization, membership=request.membership)
    payment_form = checkout_forms[1] if checkout_forms else CartPaymentForm(cart_id=cart.pk)
    if request.method == 'POST' and checkout_forms is None:
        if cart.status != Cart.Status.DRAFT:
            messages.error(request, 'Only draft carts can be edited.')
            return redirect('inventory:cart_detail', pk=cart.pk)
        if cart_form.is_valid():
            try:
                with transaction.atomic():
                    cart_form.save()
                    CartService.refresh_cart_prices(cart=cart)
                    # Validate the persisted, repriced cart before committing.
                    _cart_workspace(cart)
                    CartService.validate_discount_authority(
                        cart=cart, actor=request.user, membership=request.membership,
                    )
            except (BillingServiceError, InventoryError) as exc:
                cart.refresh_from_db()
                cart_form.add_error('discount_amount', str(exc))
            else:
                if _is_pos_request(request):
                    return _pos_response(request, cart, message='Sale details saved.')
                messages.success(request, 'Cart draft saved. Stock remains unchanged.')
                return redirect('inventory:cart_detail', pk=cart.pk)
        if _is_pos_request(request):
            # ModelForm validation mutates its in-memory instance even when it
            # is invalid. Reload before returning authoritative saved totals.
            cart.refresh_from_db()
            return _pos_response(
                request,
                cart,
                message=_form_error_summary(cart_form) or 'Check the sale details and try again.',
                level='warning',
                status=422,
                errors=cart_form.errors.get_json_data(),
            )
    workspace = _cart_workspace(cart)
    lines = workspace['lines']
    query = request.GET.get('q', '').strip()
    category_id = request.GET.get('category', '').strip()
    existing_line = CartLine.objects.filter(cart=cart, product_id=OuterRef('pk'))
    cost_movement = StockMovement.objects.filter(tenant=organization, product_id=OuterRef('pk'))
    catalog = Product.objects.filter(tenant=organization, is_active=True).select_related(
        'catalog_category', 'inventory_balance'
    ).annotate(
        in_cart=Exists(existing_line),
        cart_quantity=Subquery(existing_line.values('quantity')[:1]),
        has_cost_movement=Exists(cost_movement),
    ).order_by('name')
    if query:
        catalog = catalog.filter(Q(name__icontains=query) | Q(sku__icontains=query) | Q(brand__icontains=query) | Q(model_number__icontains=query))
    if category_id.isdigit():
        catalog = catalog.filter(catalog_category_id=category_id)
    catalog = list(catalog[:50])
    for product in catalog:
        intended_quantity = (product.cart_quantity or Decimal('0')) + Decimal('1')
        product.pos_price, product.pos_pricing_mode = CartService.line_pricing(
            product=product, quantity=intended_quantity, customer=cart.customer,
            sale_pricing_category=cart.sale_pricing_category,
        )
        balance = getattr(product, 'inventory_balance', None)
        floor = (
            Decimal(balance.average_cost)
            if balance is not None and balance.quantity > 0 and product.has_cost_movement
            else Decimal(product.buying_price or 0)
        )
        product.pos_ready = product.pos_price > floor
        product.pos_readiness_message = (
            '' if product.pos_ready else
            'Pricing review required before this product can be added to a sale.'
        )
    return render(request, 'inventory/cart_detail.html', {
        'cart': cart, 'cart_form': cart_form, 'payment_form': payment_form, **workspace,
        'is_stale': stale(cart),
        'can_restore': cart.status == Cart.Status.ABANDONED and not cart.converted_at and not business_linked(cart) and has_tenant_permission(request.user, organization, PermissionCode.FINANCE_SALES_VIEW_ALL, membership=request.membership),
        'can_apply_line_discount': has_tenant_permission(
            request.user, organization, PermissionCode.CART_DISCOUNT_APPLY, membership=request.membership,
        ),
        'catalog': catalog, 'query': query, 'category_id': category_id,
        'categories': ProductCategory.objects.filter(tenant=organization, is_active=True).order_by('name'),
        'can_register_payment': has_tenant_permission(request.user, organization, PermissionCode.PAYMENT_REGISTER, membership=request.membership),
    })


@login_required
@cart_access
def cart_line_form(request, cart_pk, line_pk=None):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    cart = get_object_or_404(Cart.objects.filter(tenant=organization, status=Cart.Status.DRAFT), pk=cart_pk)
    line = get_object_or_404(CartLine, pk=line_pk, cart=cart) if line_pk else None
    initial = {'product': request.GET.get('product')} if not line and request.GET.get('product') else None
    form = CartLineForm(
        request.POST or None, instance=line, organization=organization, initial=initial,
        sale_pricing_category=cart.sale_pricing_category, cart=cart, membership=request.membership,
    )
    selected_product_id = request.POST.get('product') if request.method == 'POST' else (
        request.GET.get('product') or (line.product_id if line else None)
    )
    selected_product = Product.objects.filter(tenant=organization, pk=selected_product_id).first() if selected_product_id else None
    if request.method == 'POST' and form.is_valid():
        unit_price, _ = CartService.line_pricing(
            product=form.cleaned_data['product'], quantity=form.cleaned_data['quantity'], customer=cart.customer,
            sale_pricing_category=cart.sale_pricing_category,
        )
        if form.cleaned_data['discount_amount'] > form.cleaned_data['quantity'] * unit_price:
            form.add_error('discount_amount', 'Discount cannot exceed the effective sale price.')
            return render(request, 'inventory/cart_line_form.html', {
                'form': form, 'cart': cart, 'line': line, 'selected_product': selected_product,
                'title': 'Edit cart item' if line else 'Add cart item',
            })
        try:
            CartService.save_line(
                organization=organization, cart_id=cart.pk, data=request.POST,
                actor=request.user, membership=request.membership, line_id=line_pk,
            )
        except ValidationError as exc:
            if hasattr(exc, 'message_dict'):
                for field, errors in exc.message_dict.items():
                    form.add_error(field if field in form.fields else None, errors)
            else:
                form.add_error(None, exc)
            return render(request, 'inventory/cart_line_form.html', {
                'form': form, 'cart': cart, 'line': line, 'selected_product': selected_product,
                'title': 'Edit cart item' if line else 'Add cart item',
            })
        except InventoryError as exc:
            form.add_error('discount_amount', str(exc))
            return render(request, 'inventory/cart_line_form.html', {
                'form': form, 'cart': cart, 'line': line, 'selected_product': selected_product,
                'title': 'Edit cart item' if line else 'Add cart item',
            })
        messages.success(request, 'Cart item saved. Stock remains unchanged until full payment.')
        return redirect('inventory:cart_detail', pk=cart.pk)
    return render(request, 'inventory/cart_line_form.html', {
        'form': form, 'cart': cart, 'line': line, 'selected_product': selected_product,
        'title': 'Edit cart item' if line else 'Add cart item',
    })


@login_required
@require_POST
@cart_access
def cart_line_discount(request, cart_pk, line_pk):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    require_permission(request, PermissionCode.CART_DISCOUNT_APPLY)
    cart = get_object_or_404(
        Cart.objects.filter(tenant=organization, status=Cart.Status.DRAFT), pk=cart_pk,
    )
    get_object_or_404(CartLine.objects.filter(tenant=organization, cart=cart), pk=line_pk)
    form = CartLineDiscountForm(request.POST)
    if not form.is_valid():
        message = _form_error_summary(form)
        if _is_pos_request(request):
            return _pos_response(request, cart, message=message, level='warning', status=422,
                                 errors=form.errors.get_json_data())
        messages.error(request, message)
        return redirect('inventory:cart_detail', pk=cart.pk)
    try:
        cart = CartService.set_line_discount(
            organization=organization, cart_id=cart.pk, line_id=line_pk,
            discount_amount=form.cleaned_data['discount_amount'], actor=request.user,
            membership=request.membership,
        )
    except InventoryError as exc:
        if _is_pos_request(request):
            cart.refresh_from_db()
            return _pos_response(request, cart, message=str(exc), level='warning', status=422,
                                 errors={'discount_amount': [{'message': str(exc), 'code': ''}]})
        messages.error(request, str(exc))
        return redirect('inventory:cart_detail', pk=cart.pk)
    if _is_pos_request(request):
        return _pos_response(request, cart, message='Item discount saved.')
    messages.success(request, 'Item discount saved.')
    return redirect('inventory:cart_detail', pk=cart.pk)


@login_required
@cart_access
def cart_line_adjust(request, cart_pk):
    """Fast POS quantity controls for standard, non-serialized catalog items."""
    organization = _scope(request, PermissionCode.CART_MANAGE)
    if request.method != 'POST':
        raise Http404
    product_id = request.POST.get('product', '')
    direction = request.POST.get('direction', 'add')
    if direction not in {'add', 'increase', 'decrease'} or not product_id.isdigit():
        raise Http404

    with transaction.atomic():
        cart = get_object_or_404(
            Cart.objects.select_for_update().filter(tenant=organization, status=Cart.Status.DRAFT), pk=cart_pk
        )
        product = get_object_or_404(Product.objects.select_for_update().filter(tenant=organization, is_active=True), pk=product_id)
        if product.is_serialized:
            if _is_pos_request(request):
                return JsonResponse({
                    'ok': False, 'message': 'This item needs serial-number selection before it can be added.',
                    'redirect_url': f'/inventory/carts/{cart.pk}/items/new/?product={product.pk}',
                }, status=422)
            messages.info(request, 'Select serial numbers before adding this serialized product.')
            return redirect('inventory:cart_line_create', cart_pk=cart.pk)
        line = CartLine.objects.select_for_update().filter(cart=cart, product=product).first()
        if direction == 'decrease':
            if line is None:
                return redirect('inventory:cart_detail', pk=cart.pk)
            if line.quantity <= Decimal('1.00'):
                line.delete()
            else:
                line.quantity -= Decimal('1.00')
                line.unit_price, _ = CartService.line_pricing(product=product, quantity=line.quantity, customer=cart.customer, sale_pricing_category=cart.sale_pricing_category)
                from products.pricing import BelowCostError, validate_net_unit_price
                try:
                    validate_net_unit_price(product=product, quantity=line.quantity, unit_price=line.unit_price,
                                            line_discount=line.discount_amount, actor=request.user, organization=organization)
                except BelowCostError as exc:
                    if _is_pos_request(request):
                        return _pos_response(request, cart, message=str(exc), level='error', status=422)
                    messages.error(request, str(exc))
                    return redirect('inventory:cart_detail', pk=cart.pk)
                line.save(update_fields=['quantity', 'unit_price', 'updated_at'])
        else:
            quantity = (line.quantity if line else Decimal('0.00')) + Decimal('1.00')
            if product.item_type == Product.ItemType.PHYSICAL and product.track_stock and quantity > product.available_stock:
                if _is_pos_request(request):
                    return _pos_response(request, cart, message=f'Only {format_quantity(product.available_stock)} {product.measure_unit} of {product.name} are available.', level='warning', status=409)
                messages.warning(request, f'Only {format_quantity(product.available_stock)} {product.measure_unit} of {product.name} are available.')
            elif line:
                line.quantity = quantity
                line.unit_price, _ = CartService.line_pricing(product=product, quantity=quantity, customer=cart.customer, sale_pricing_category=cart.sale_pricing_category)
                from products.pricing import BelowCostError, validate_net_unit_price
                try:
                    validate_net_unit_price(product=product, quantity=quantity, unit_price=line.unit_price,
                                            line_discount=line.discount_amount, actor=request.user, organization=organization)
                except BelowCostError as exc:
                    if _is_pos_request(request):
                        return _pos_response(request, cart, message=str(exc), level='error', status=422)
                    messages.error(request, str(exc))
                    return redirect('inventory:cart_detail', pk=cart.pk)
                line.save(update_fields=['quantity', 'unit_price', 'updated_at'])
            else:
                unit_price, _ = CartService.line_pricing(product=product, quantity=quantity, customer=cart.customer, sale_pricing_category=cart.sale_pricing_category)
                from products.pricing import BelowCostError, validate_net_unit_price
                try:
                    validate_net_unit_price(product=product, quantity=quantity, unit_price=unit_price,
                                            actor=request.user, organization=organization)
                except BelowCostError as exc:
                    if _is_pos_request(request):
                        return _pos_response(request, cart, message=str(exc), level='error', status=422)
                    messages.error(request, str(exc))
                    return redirect('inventory:cart_detail', pk=cart.pk)
                CartLine.objects.create(cart=cart, product=product, quantity=quantity, unit_price=unit_price)
    if _is_pos_request(request):
        return _pos_response(request, cart)
    return redirect('inventory:cart_detail', pk=cart.pk)


@login_required
@cart_access
def cart_line_delete(request, cart_pk, line_pk):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    if request.method != 'POST':
        raise Http404
    with transaction.atomic():
        cart = get_object_or_404(
            Cart.objects.select_for_update().filter(
                tenant=organization, status=Cart.Status.DRAFT,
            ),
            pk=cart_pk,
        )
        line = get_object_or_404(
            CartLine.objects.select_for_update().select_related('product'),
            pk=line_pk,
            cart=cart,
            tenant=organization,
        )
        serial_unit_ids = list(
            line.serial_selections.order_by('stock_unit_id').values_list(
                'stock_unit_id', flat=True,
            )
        )
        audit(
            organization=organization,
            actor=request.user,
            action='inventory.cart_line.removed',
            obj=line,
            old_value={
                'cart_id': cart.pk,
                'product_id': line.product_id,
                'product_name': line.product.name,
                'quantity': str(line.quantity),
                'unit_price': str(line.unit_price),
                'discount_amount': str(line.discount_amount),
                'serial_unit_ids': serial_unit_ids,
            },
            metadata={
                'cart_status': cart.status,
                'serialized': line.product.is_serialized,
                'stock_changed': False,
            },
        )
        line.delete()
    if _is_pos_request(request):
        return _pos_response(
            request, cart, message='Item removed from the draft sale. Stock was not changed.',
        )
    messages.success(request, 'Item removed from the draft sale. Stock was not changed.')
    return redirect('inventory:cart_detail', pk=cart.pk)


@login_required
@cart_access
def cart_abandon(request, pk):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    if request.method != 'POST':
        raise Http404
    get_object_or_404(Cart.objects.filter(tenant=organization), pk=pk)
    try:
        _cart, changed = CartService.abandon(
            organization=organization,
            cart_id=pk,
            actor=request.user,
            reason=request.POST.get('reason', ''),
            membership=request.membership,
        )
    except InventoryError as exc:
        messages.error(request, str(exc))
    else:
        if changed:
            messages.success(request, 'Sale discarded. Its audit record was preserved and no stock was changed.')
        else:
            messages.info(request, 'This sale was already discarded.')
    return redirect('inventory:cart_list')


@login_required
@cart_access
def cart_convert(request, pk, target):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    if request.method != 'POST':
        raise Http404
    permission = PermissionCode.BILLING_CREATE
    require_permission(request, permission)
    try:
        with transaction.atomic():
            cart = get_object_or_404(
                Cart.objects.select_for_update().filter(tenant=organization), pk=pk,
            )
            # Modern checkout buttons belong to the details form, so the exact
            # customer and financial values the user confirms arrive here. Keep
            # accepting legacy CSRF-only posts for backward compatibility.
            if cart.status == Cart.Status.DRAFT and 'discount_amount' in request.POST:
                cart_form = CartForm(
                    request.POST,
                    instance=cart,
                    organization=organization,
                    membership=request.membership,
                )
                if not cart_form.is_valid():
                    raise InventoryError('Sale was not converted. ' + _form_error_summary(cart_form))
                cart_form.save()
                CartService.validate_restored(cart=cart)
                CartService.refresh_cart_prices(cart=cart)
                _cart_workspace(cart)
                CartService.validate_discount_authority(
                    cart=cart, actor=request.user, membership=request.membership,
                )
            document = CartService.convert(
                organization=organization, cart_id=pk, target=target, actor=request.user, membership=request.membership,
            )
        messages.success(request, f'Cart converted to {target} {document.number}.')
        if target == BillingDocument.DocumentType.INVOICE and has_tenant_permission(request.user, organization, PermissionCode.PAYMENT_REGISTER, membership=request.membership):
            return redirect('billing:create_receipt_from_invoice', pk=document.pk)
        return redirect('billing:document_detail', doc_type=target, pk=document.pk)
    except (InventoryError, BillingServiceError) as exc:
        messages.error(request, str(exc))
        return redirect('inventory:cart_detail', pk=pk)


@login_required
@require_POST
@cart_access
def cart_checkout(request, pk):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    require_permission(request, PermissionCode.PAYMENT_REGISTER)
    payment_form = CartPaymentForm(request.POST, cart_id=pk)
    try:
        with transaction.atomic():
            cart = get_object_or_404(Cart.objects.select_for_update().filter(tenant=organization), pk=pk)
            cart_form = CartForm(
                request.POST, instance=cart, organization=organization, membership=request.membership,
            )
            if cart.status == Cart.Status.DRAFT:
                details_valid = cart_form.is_valid()
                payment_valid = payment_form.is_valid()
                if not details_valid or not payment_valid:
                    raise InventoryError('Check the sale and payment details below.')
                cart_form.save()
                CartService.refresh_cart_prices(cart=cart)
                _cart_workspace(cart)
            receipt = CartService.checkout(
                organization=organization, cart_id=pk, actor=request.user, membership=request.membership,
                payment_method=request.POST.get('payment_method', '').strip(),
                payment_reference=request.POST.get('payment_reference', '').strip(),
            )
    except (InventoryError, BillingServiceError) as exc:
        cart_form.add_error(None, str(exc))
        response = cart_detail(request, pk=pk, checkout_forms=(cart_form, payment_form))
        response.status_code = 422
        return response
    messages.success(request, 'Payment recorded. Sale completed.')
    return redirect('billing:document_detail', doc_type='receipt', pk=receipt.pk)


@login_required
@require_POST
@cart_access
def cart_restore(request, pk):
    organization = _scope(request, PermissionCode.CART_MANAGE)
    try:
        CartService.restore(organization=organization, cart_id=pk, actor=request.user, membership=request.membership)
        messages.success(request, 'Cart restored. Review current product eligibility, prices, stock and serials before conversion; saved quantities and prices were preserved.')
    except InventoryError as exc:
        messages.error(request, str(exc))
    return redirect('inventory:cart_detail', pk=pk)


REPORT_NAMES = {
    'stock-valuation': 'Stock valuation', 'purchases': 'Purchase report', 'sales': 'Sales report',
    'gross-profit': 'Gross profit', 'best-selling': 'Best-selling products', 'fast-moving': 'Fast-moving products',
    'slow-moving': 'Slow-moving products', 'dead-stock': 'Dead stock', 'low-stock': 'Low-stock products',
    'out-of-stock': 'Out-of-stock products', 'movements': 'Stock movements', 'serialized': 'Serialized stock history',
    'historical-purchases': 'Imported purchases',
    'historical-sales': 'Imported sales',
}
COST_REPORT_NAMES = {'stock-valuation', 'purchases', 'gross-profit', 'movements', 'historical-purchases'}

REPORT_DESCRIPTIONS = {
    'stock-valuation': 'Current movement-backed quantity and recorded acquisition value.',
    'purchases': 'Confirmed stock receipts in the selected period.',
    'sales': 'Fully paid inventory sales with stock already deducted.',
    'gross-profit': 'Operational gross profit, not accounting net profit. Stock costs are TZS; profit for non-TZS revenue is unavailable without an exchange-rate snapshot.',
    'best-selling': 'Products ranked by units sold in the configured fast-moving period.',
    'fast-moving': 'Products meeting the configured high-sales threshold.',
    'slow-moving': 'Products at or below the configured low-sales threshold.',
    'dead-stock': 'Stock on hand with no sale movement during the configured dead-stock period.',
    'low-stock': 'Products at or below their reorder threshold but still available.',
    'out-of-stock': 'Stock-tracked products with no available units.',
    'movements': 'Immutable stock in, sale, adjustment, and opening-balance records.',
    'serialized': 'Lifecycle history for individually tracked serial numbers.',
    'historical-purchases': 'Imported reference-only purchase history; it does not affect live stock.',
    'historical-sales': 'Imported reference-only sales history; it does not affect live stock.',
}




@login_required
def report(request, report_name='stock-valuation'):
    if report_name not in REPORT_NAMES:
        raise Http404
    sensitive = report_name in COST_REPORT_NAMES
    organization = _scope(request, PermissionCode.COST_REPORT_VIEW if sensitive else PermissionCode.SALES_REPORT_VIEW)
    exporting = request.GET.get('export') == 'xlsx'
    if exporting:
        require_permission(request, PermissionCode.INVENTORY_EXPORT)
    form = ReportFilterForm(request.GET, organization=organization, report_name=report_name)
    valid = form.is_valid()
    filters = form.cleaned_data if valid else {}
    headers, rows = report_rows(organization, report_name, filters) if valid else ([], [])
    rows = filter_report_rows(report_name, rows, filters) if valid else []
    bridge = valuation_bridge(organization, rows) if valid and report_name == 'stock-valuation' else None
    totals = report_totals(report_name, rows) if valid else []
    summaries = [(label, total) for label, column, total in totals]
    date_from = request.GET.get('date_from', '') if report_name in DATED_REPORTS else ''
    date_to = request.GET.get('date_to', '') if report_name in DATED_REPORTS else ''
    params = request.GET.copy()
    params.pop('page', None)
    params.pop('export', None)
    params.pop('search_in', None)
    filter_query = params.urlencode()
    if exporting and valid:
        # Excel cannot store timezone-aware datetimes. Export local wall time,
        # and force text cells to prevent supplier/SKU formula injection.
        export_rows = [list(row) for row in rows]
        for label, column, total in totals:
            total_row = [''] * len(headers)
            total_row[0] = label
            total_row[column] = total
            export_rows.append(total_row)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = REPORT_NAMES[report_name][:31]
        for row in [headers, *export_rows]:
            sheet.append([value.replace(tzinfo=None) if hasattr(value, 'tzinfo') and value.tzinfo else value for value in row])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        metadata = workbook.create_sheet('Report context')
        metadata.append(['Report', REPORT_NAMES[report_name]])
        metadata.append(['Timezone', timezone.get_current_timezone_name()])
        for key, value in filters.items():
            metadata.append([key, str(value) if value is not None else 'All'])
            for cell in metadata[metadata.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        if bridge is not None:
            breakdown = workbook.create_sheet('Valuation bridge')
            breakdown.append(['Source (all matching products, all time)', 'Value (TZS)'])
            for label, value in bridge['entries']:
                breakdown.append([label, value])
            breakdown.append(['Current stock value', bridge['total']])
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{report_name}.xlsx"'
        response['Cache-Control'] = 'private, no-store'
        return response
    page_obj = Paginator(rows, 50).get_page(request.GET.get('page'))
    display_headers, display_rows = report_display(headers, page_obj.object_list)
    loss_count = sum(1 for row in rows if row[7] is not None and row[7] < 0) if report_name == 'gross-profit' else 0
    allowed_reports = {
        slug: label for slug, label in REPORT_NAMES.items()
        if has_tenant_permission(request.user, organization,
            PermissionCode.COST_REPORT_VIEW if slug in COST_REPORT_NAMES else PermissionCode.SALES_REPORT_VIEW,
            membership=request.membership)
    }
    report_groups = []
    for label, slugs in (
        ('Stock & inventory', ('stock-valuation', 'low-stock', 'out-of-stock', 'dead-stock', 'movements', 'serialized')),
        ('Purchasing', ('purchases',)),
        ('Sales & margin', ('sales', 'gross-profit', 'best-selling', 'fast-moving', 'slow-moving')),
        ('Imported records', ('historical-purchases', 'historical-sales')),
    ):
        links = [(slug, allowed_reports[slug]) for slug in slugs if slug in allowed_reports]
        if links:
            report_groups.append({'label': label, 'links': links})
    response = render(request, 'inventory/report.html', {
        'display_headers': display_headers, 'display_rows': display_rows, 'valuation_bridge': bridge,
        'loss_count': loss_count, 'is_imported_report': report_name.startswith('historical-'),
        'report_name': report_name,
        'title': REPORT_NAMES[report_name],
        'report_names': allowed_reports,
        'report_groups': report_groups,
        'description': REPORT_DESCRIPTIONS[report_name], 'result_count': len(rows),
        'headers': headers, 'rows': page_obj.object_list, 'page_obj': page_obj,
        'query': request.GET.get('q', ''), 'date_from': date_from, 'date_to': date_to,
        'filter_form': form, 'summaries': summaries, 'filter_query': filter_query,
        'filters_valid': valid, 'dated_report': report_name in DATED_REPORTS,
    })
    response.status_code = 200 if valid else 400
    response['Cache-Control'] = 'private, no-store'
    return response


@login_required
def export_records(request, record_type):
    organization = _scope(request, PermissionCode.INVENTORY_EXPORT)
    if record_type == 'products':
        headers = ['SKU', 'Name', 'Type', 'Category', 'Brand', 'Model', 'Selling price', 'Technician price', 'Wholesale price', 'Active']
        rows = [[
            p.sku, p.name, p.get_item_type_display(),
            p.catalog_category.name if p.catalog_category_id else p.get_category_display(),
            p.brand, p.model_number, p.selling_price, p.technician_price, p.wholesale_price, p.is_active,
        ] for p in Product.objects.filter(tenant=organization).select_related('catalog_category').order_by('name')]
    elif record_type == 'suppliers':
        require_permission(request, PermissionCode.SUPPLIER_MANAGE)
        headers = ['Company', 'Contact', 'Phone', 'Email', 'Address', 'TIN/VRN', 'Active']
        rows = [[s.company_name, s.contact_person, s.phone, s.email, s.physical_address, s.tin_vrn, s.is_active] for s in Supplier.objects.filter(tenant=organization).order_by('company_name')]
    else:
        raise Http404
    response = HttpResponse(workbook_bytes(headers, rows, title=record_type.title()), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{record_type}.xlsx"'
    return response


@login_required
def invoice_serials(request, pk):
    organization = _scope(request, PermissionCode.BILLING_CREATE)
    invoice = get_object_or_404(
        BillingDocument.objects.filter(
            tenant=organization, document_type=BillingDocument.DocumentType.INVOICE,
            status__in=[BillingDocument.Status.DRAFT, BillingDocument.Status.ISSUED],
        ).prefetch_related('items__product'),
        pk=pk,
    )
    lines = list(invoice.items.filter(product__is_serialized=True).select_related('product').order_by('id'))
    line_options = []
    for line in lines:
        selected_ids = set(DocumentSerialSelection.objects.filter(tenant=organization, billing_line=line).values_list('stock_unit_id', flat=True))
        units = StockUnit.objects.filter(
            tenant=organization, product=line.product, status=StockUnit.Status.AVAILABLE
        ).order_by('serial_number')
        line_options.append({'line': line, 'units': units, 'selected_ids': selected_ids})
    if request.method == 'POST':
        errors = []
        selections = {}
        seen_serial_ids = set()
        for entry in line_options:
            line = entry['line']
            ids = request.POST.getlist(f'line_{line.pk}')
            normalized_ids = {int(value) for value in ids
                              if value.isdecimal() and len(value) <= 19 and 0 < int(value) <= 9223372036854775807}
            entry['selected_ids'] = normalized_ids
            if any(not value.isdecimal() or len(value) > 19 or not 0 < int(value) <= 9223372036854775807 for value in ids):
                errors.append(f'A selected serial for {line.product.name} is unavailable.')
                continue
            if len(normalized_ids) != len(ids) or seen_serial_ids.intersection(normalized_ids):
                errors.append('Each serial may be selected only once across this invoice.')
            seen_serial_ids.update(normalized_ids)
            if len(ids) != int(line.quantity):
                errors.append(f'{line.product.name} requires exactly {int(line.quantity)} serial numbers.')
                continue
            units = list(StockUnit.objects.filter(
                tenant=organization, product=line.product, status=StockUnit.Status.AVAILABLE, pk__in=normalized_ids
            ))
            if len(units) != len(ids):
                errors.append(f'A selected serial for {line.product.name} is unavailable.')
            selections[line.pk] = units
        if not errors:
            with transaction.atomic():
                DocumentSerialSelection.objects.filter(tenant=organization, billing_line__in=lines).delete()
                DocumentSerialSelection.objects.bulk_create([
                    DocumentSerialSelection(
                        organization=organization, tenant=organization, billing_line_id=line_id, stock_unit=unit
                    ) for line_id, units in selections.items() for unit in units
                ])
                audit(organization=organization, actor=request.user, action='inventory.invoice.serials_selected', obj=invoice)
            messages.success(request, 'Serial selections saved. Availability will be checked again at full payment.')
            return redirect('billing:document_detail', doc_type='invoice', pk=invoice.pk)
        for error in errors:
            messages.error(request, error)
    return render(request, 'inventory/invoice_serials.html', {'invoice': invoice, 'line_options': line_options})


@login_required
def import_data(request):
    organization = _scope(request, PermissionCode.INVENTORY_IMPORT)
    form = InventoryImportForm(request.POST or None, request.FILES or None, organization=organization)
    job = None
    if request.method == 'POST' and form.is_valid():
        job = validate_workbook(
            organization=organization, actor=request.user, import_type=form.cleaned_data['import_type'], uploaded_file=form.cleaned_data['workbook']
        )
        if job.error_count:
            messages.error(request, 'Validation found errors. Nothing was imported.')
        else:
            messages.success(request, 'Validation succeeded. Review and commit the import.')
    return render(request, 'inventory/import.html', {'form': form, 'job': job, 'recent_jobs': ImportJob.objects.filter(tenant=organization)[:10]})


@login_required
def import_commit(request, pk):
    organization = _scope(request, PermissionCode.INVENTORY_IMPORT)
    if request.method != 'POST':
        raise Http404
    try:
        job = commit_import(organization=organization, actor=request.user, job_id=pk)
        messages.success(request, f'Imported {job.row_count} rows successfully.')
    except (ValueError, InventoryError) as exc:
        messages.error(request, str(exc))
    return redirect('inventory:import_data')


@login_required
def import_template(request, import_type):
    _scope(request, PermissionCode.INVENTORY_IMPORT)
    try:
        payload = template_workbook(import_type)
    except ValueError:
        raise Http404
    response = HttpResponse(payload, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{import_type}-template.xlsx"'
    return response


@login_required
def settings_edit(request):
    organization = _scope(request, PermissionCode.CATEGORY_MANAGE)
    settings_obj, _ = InventorySettings.objects.get_or_create(organization=organization, tenant=organization)
    form = InventorySettingsForm(request.POST or None, instance=settings_obj, organization=organization)
    if request.method == 'POST' and form.is_valid():
        obj = form.save(commit=False)
        obj.organization = obj.tenant = organization
        obj.updated_by = request.user
        obj.save()
        audit(organization=organization, actor=request.user, action='inventory.settings.updated', obj=obj)
        messages.success(request, 'Inventory settings saved.')
        return redirect('inventory:dashboard')
    return render(request, 'inventory/form.html', {
        'form': form, 'title': 'Inventory settings',
        'subtitle': 'Configure operational thresholds used by dashboard alerts and movement reports.',
        'cancel_url': 'inventory:dashboard', 'submit_label': 'Save settings',
    })
