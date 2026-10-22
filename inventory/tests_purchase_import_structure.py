"""Synthetic clipboard fixtures: no customer or supplier data."""
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl import Workbook

from .purchase_import import ImportInputError, parse_input


class PurchaseImportStructureTests(SimpleTestCase):
    def test_concatenated_product_records_fail_before_resolution(self):
        # Plausible reproduction of a lost-delimiter clipboard, not the original.
        for text in ('Dell Latitude 5420 2 650000 TP-Link Archer C6 10 85000',
                     'LAPTOPDELL Latitude 5420LAP00151,200,000LAPTOPHP ProBook G9LAP00231350,000'):
            with self.subTest(text=text), self.assertRaisesMessage(ImportInputError, 'spreadsheet columns'):
                parse_input(text=text)
        with self.assertRaisesMessage(ImportInputError, 'spreadsheet columns'):
            parse_input(text='LAPTOPDELL Latitude 5420LAP00151,200,000', input_format='mixed')

    def test_headers_map_by_name_preserving_blank_sku_and_intermediate_cells(self):
        rows = parse_input(text='Product name\tSKU\tUnit cost\tQuantity\tBrand\tModel\n'
                                'Dell Latitude 5420\t\t650000\t2\t\t5420')
        self.assertEqual(rows[0]['product_name'], 'Dell Latitude 5420')
        self.assertEqual(rows[0]['sku'], '')
        self.assertEqual(rows[0]['quantity'], '2')
        self.assertEqual(rows[0]['brand'], '')
        self.assertEqual(rows[0]['model_number'], '5420')

    def test_blank_lines_and_embedded_newlines_preserve_source_numbers(self):
        for newline in ('\n', '\r\n'):
            rows = parse_input(text=newline.join([
                'SKU\tProduct name\tQuantity\tUnit cost', '',
                '0001\t"Router\tA' + newline + 'radio"\t1\t10', '',
                '0002\tCable\t2\t3',
            ]))
            self.assertEqual([r['row'] for r in rows], [3, 6])
            self.assertEqual(rows[0]['sku'], '0001')
            self.assertIn('Router\tA', rows[0]['product_name'])

    def test_bad_headers_and_extra_columns_fail_safely(self):
        for text in ('sku\tquantity\tqty\tunit_cost\nA\t1\t1\t2',
                     'sku\tproduct_name\tquantity\twrong_cost\nA\tRadio\t1\t2',
                     'A\t1\t2\t\t\t\textra',
                     r'A\t1\t2\nB\t1\t2'):
            with self.subTest(text=text), self.assertRaises(ImportInputError):
                parse_input(text=text)

    def test_headerless_mixed_input_requires_format(self):
        text = '\tDell Latitude\t2\t650000'
        with self.assertRaisesMessage(ImportInputError, 'format'):
            parse_input(text=text)
        self.assertEqual(parse_input(text=text, input_format='mixed')[0]['quantity'], '2')

    def test_upload_and_paste_equivalence_with_alias_headers(self):
        cells = [('Product name', 'SKU', 'Quantity', 'Unit cost', 'Serials'),
                 ('Radio, model A', '000123', '1', '1,250.50', '00001')]
        book = Workbook()
        book.active.title = 'Instructions'
        book.active.append(['Never import this'])
        sheet = book.create_sheet('Purchase rows')
        for row in cells:
            sheet.append(row)
        output = BytesIO()
        book.save(output)
        upload = SimpleUploadedFile('synthetic.xlsx', output.getvalue())
        pasted = parse_input(text='\n'.join('\t'.join(row) for row in cells))
        self.assertEqual(parse_input(upload=upload), pasted)

    def test_formula_cells_are_row_errors_and_identifiers_stay_text(self):
        rows = parse_input(text='sku\tproduct_name\tquantity\tunit_cost\n0001\t=HYPERLINK("x")\t1\t2')
        self.assertEqual(rows[0]['sku'], '0001')
        self.assertIn('Formula', rows[0]['input_errors'][0])

    def test_numeric_grouping_is_explicit_and_does_not_change_names_or_serials(self):
        row = parse_input(text='SKU\tProduct name\tQuantity\tUnit cost\tSelling price\tSerials\n'
                               '0001\tRadio, white\t1,000\t1,250.50\t2,500.00\t00001,00002')[0]
        self.assertEqual((row['quantity'], row['unit_cost'], row['selling_price']), ('1000', '1250.50', '2500.00'))
        self.assertEqual((row['product_name'], row['serial_numbers']), ('Radio, white', '00001,00002'))
        for cost in ('1,25', '1.250,50', 'TZS 1250'):
            with self.subTest(cost=cost), self.assertRaisesMessage(ImportInputError, 'Row 2, unit cost'):
                parse_input(text='SKU\tQuantity\tUnit cost\n0001\t1\t' + cost)

    def test_worksheet_headers_and_formula_policy(self):
        book = Workbook()
        book.active.title = 'Purchase rows'
        book.active.append(['sku', 'quantity', 'unit_cost'])
        book.active.append(['0001', 1, '=1+1'])
        output = BytesIO()
        book.save(output)
        rows = parse_input(upload=SimpleUploadedFile('synthetic.xlsx', output.getvalue()))
        self.assertEqual(rows[0]['row'], 2)
        self.assertIn('Formula', rows[0]['input_errors'][0])
        book.active.delete_rows(1)
        output = BytesIO()
        book.save(output)
        with self.assertRaisesMessage(ImportInputError, 'headers'):
            parse_input(upload=SimpleUploadedFile('synthetic.xlsx', output.getvalue()))

    def test_exactly_two_hundred_nonempty_rows_and_original_locations(self):
        rows = parse_input(text='sku\tquantity\tunit_cost\n' + '\n\n'.join('0001\t1\t2' for _ in range(200)))
        self.assertEqual(len(rows), 200)
        self.assertEqual(rows[-1]['row'], 400)
        book = Workbook()
        book.active.title = 'Purchase rows'
        book.active.append(['sku', 'quantity', 'unit_cost'])
        for number in range(2, 401, 2):
            for column, value in enumerate(('0001', '1', '2'), 1):
                book.active.cell(number, column, value)
        output = BytesIO()
        book.save(output)
        self.assertEqual(parse_input(upload=SimpleUploadedFile('synthetic.xlsx', output.getvalue())), rows)
