from datetime import date, datetime, timezone as dt_timezone
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from products.models import Product
from users.models import Organization, UserAccessProfile
from .models import (InventoryBalance, InventorySettings, Purchase, PurchaseLine,
                     StockMovement, Supplier, InventorySale, InventorySaleLine, HistoricalInventoryRecord, ImportJob, StockUnit)
from .reporting import report_rows
from .views import REPORT_NAMES
from billing.models import BillingDocument, BillingLineItem
from customers.models import Customer


class InventoryReportTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name='Reports', slug='reports')
        self.other = Organization.objects.create(name='Other reports', slug='other-reports')
        self.admin = get_user_model().objects.create_user(username='report-admin', password='pass')
        UserAccessProfile.objects.create(user=self.admin, tenant=self.org, role=UserAccessProfile.Role.TENANT_ADMIN)
        self.staff = get_user_model().objects.create_user(username='report-staff', password='pass')
        UserAccessProfile.objects.create(user=self.staff, tenant=self.org, role=UserAccessProfile.Role.TENANT_STAFF)
        self.supplier = Supplier.objects.create(tenant=self.org, company_name='=Supplier', is_active=False)
        self.other_supplier = Supplier.objects.create(tenant=self.other, company_name='Private supplier')
        self.product = Product.objects.create(tenant=self.org, organization=self.org, sku='REPORT-1', name='Router', measure_unit='Unit', buying_price=10, selling_price=20, retail_price=20, quantity=0, stock=0)
        self.client.force_login(self.admin)

    def url(self, name='purchases'):
        return reverse('inventory:report', args=[name])

    def purchase(self, ref='P-1', status=Purchase.Status.CONFIRMED, when=date(2026, 10, 1), supplier=None, count=1, snapshot=None):
        supplier = supplier or self.supplier
        purchase = Purchase.objects.create(tenant=supplier.tenant, supplier=supplier, reference_number=ref, purchase_date=when, created_by=self.admin)
        if supplier.tenant_id == self.org.pk:
            PurchaseLine.objects.bulk_create([
                PurchaseLine(tenant=self.org, purchase=purchase, product=self.product, quantity=Decimal('2'), unit_cost=Decimal('10.123456'), authoritative_purchase_total=snapshot)
                for _ in range(count)
            ])
        # Fixture construction only; production confirmation remains service-owned.
        Purchase.objects.filter(tenant=supplier.tenant, pk=purchase.pk).update(status=status)
        return purchase

    def test_purchase_total_filters_supplier_dates_and_status_before_pagination(self):
        self.purchase(count=51)
        self.purchase('DRAFT', status=Purchase.Status.DRAFT)
        self.purchase('CANCEL', status=Purchase.Status.CANCELLED)
        self.purchase('OLD', when=date(2026, 9, 30))
        self.purchase('PRIVATE', supplier=self.other_supplier)
        another = Supplier.objects.create(tenant=self.org, company_name='Second supplier')
        self.purchase('SECOND', supplier=another)
        params = {'supplier': self.supplier.pk, 'date_from': '2026-10-01', 'date_to': '2026-10-01'}
        response = self.client.get(self.url(), params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['result_count'], 51)
        self.assertEqual(response.context['summaries'], [('Total purchases (TZS)', Decimal('1032.75'))])
        self.assertEqual(len(response.context['rows']), 50)
        self.assertContains(response, f'supplier={self.supplier.pk}')
        self.assertNotContains(response, 'Private supplier')
        second = self.client.get(self.url(), {**params, 'page': 2})
        self.assertEqual(second.context['summaries'], response.context['summaries'])
        self.assertEqual(len(second.context['rows']), 1)

    def test_all_suppliers_total_and_snapshot_are_preserved(self):
        self.purchase(snapshot=Decimal('77.77'))
        another = Supplier.objects.create(tenant=self.org, company_name='Another')
        self.purchase('SECOND', supplier=another)
        response = self.client.get(self.url())
        self.assertEqual(response.context['summaries'][0][1], Decimal('98.02'))
        self.assertEqual(self.client.get(self.url(), {'q': 'P-1'}).context['summaries'][0][1], Decimal('77.77'))

    def test_invalid_filters_fail_closed_in_html_and_export(self):
        self.purchase()
        for params in ({'date_from': 'bad'}, {'date_to': '2026-02-30'}, {'date_from': '2026-10-02', 'date_to': '2026-10-01'}, {'supplier': self.other_supplier.pk}, {'supplier': 'bad'}):
            for export in ({}, {'export': 'xlsx'}):
                response = self.client.get(self.url(), {**params, **export})
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.context['rows'], [])
                self.assertEqual(response.context['summaries'], [])
                self.assertContains(response, 'Report not generated', status_code=400)

    def test_purchase_excel_matches_total_and_escapes_formula(self):
        self.purchase(count=51)
        response = self.client.get(self.url(), {'supplier': self.supplier.pk, 'export': 'xlsx', 'page': 2})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 53)
        self.assertEqual(sheet['C2'].value, '=Supplier')
        self.assertEqual(sheet['C2'].data_type, 's')
        self.assertEqual(sheet['A53'].value, 'Total purchases (TZS)')
        self.assertEqual(Decimal(str(sheet['H53'].value)), Decimal('1032.75'))

    def test_valuation_total_rounds_each_row_and_search_changes_scope(self):
        InventoryBalance.objects.create(tenant=self.org, product=self.product, quantity=Decimal('1.5'), average_cost=Decimal('10.123456'))
        other_product = Product.objects.create(tenant=self.other, organization=self.other, sku='SECRET', name='Private', measure_unit='Unit', buying_price=1, selling_price=2, quantity=0, stock=0)
        InventoryBalance.objects.create(tenant=self.other, product=other_product, quantity=999, average_cost=999)
        response = self.client.get(self.url('stock-valuation'))
        self.assertEqual(response.context['summaries'], [('Total stock value (TZS)', Decimal('15.19'))])
        self.assertNotContains(response, 'Private')
        empty = self.client.get(self.url('stock-valuation'), {'q': 'absent'})
        self.assertEqual(empty.context['summaries'][0][1], Decimal('0.00'))
        sheet = load_workbook(BytesIO(self.client.get(self.url('stock-valuation'), {'export': 'xlsx'}).content)).active
        self.assertEqual(Decimal(str(sheet['E3'].value)), Decimal('15.19'))
        self.assertNotIn('date_from', response.context['filter_form'].fields)

    def test_movement_local_date_and_excel_timezone(self):
        movement = StockMovement.objects.create(tenant=self.org, product=self.product, movement_type='opening_balance', quantity=1, balance_after=1, unit_cost=10)
        StockMovement.objects.filter(tenant=self.org, pk=movement.pk).update(created_at=datetime(2026, 10, 1, 22, 30, tzinfo=dt_timezone.utc))
        with timezone.override('Africa/Dar_es_Salaam'):
            params = {'date_from': '2026-10-02', 'date_to': '2026-10-02'}
            response = self.client.get(self.url('movements'), params)
            self.assertEqual(response.context['result_count'], 1)
            response = self.client.get(self.url('movements'), {**params, 'export': 'xlsx'})
            sheet = load_workbook(BytesIO(response.content)).active
            self.assertEqual(sheet['A2'].value, datetime(2026, 10, 2, 1, 30))
            self.assertEqual(self.client.get(self.url('movements'), {'date_to': '2026-10-01'}).context['result_count'], 0)

    def test_zero_sales_stock_is_slow_and_dead(self):
        InventoryBalance.objects.create(tenant=self.org, product=self.product, quantity=5, average_cost=10)
        for name in ('slow-moving', 'dead-stock'):
            headers, rows = report_rows(self.org, name)
            self.assertEqual(rows[0][0], self.product.sku)
        self.assertEqual(report_rows(self.org, 'slow-moving')[1][0][2], 0)

    def test_out_of_stock_includes_products_without_balance(self):
        self.assertEqual(report_rows(self.org, 'out-of-stock')[1][0][0], self.product.sku)

    def test_report_reads_never_create_settings(self):
        self.assertFalse(InventorySettings.objects.filter(tenant=self.org).exists())
        for name in REPORT_NAMES:
            self.assertEqual(self.client.get(self.url(name)).status_code, 200, name)
        self.assertFalse(InventorySettings.objects.filter(tenant=self.org).exists())

    def test_staff_cannot_view_or_export_financial_reports(self):
        self.client.force_login(self.staff)
        for name in REPORT_NAMES:
            for params in ({}, {'export': 'xlsx'}):
                self.assertEqual(self.client.get(self.url(name), params).status_code, 403, name)

    def test_unauthenticated_reports_redirect(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url()).status_code, 302)

    def test_sales_and_gross_profit_use_snapshots_local_date_and_separate_currencies(self):
        customer = Customer.objects.create(tenant=self.org, organization=self.org, name='Buyer', customer_type='random', location='Moshi')
        for index, currency in enumerate(('TZS', 'USD')):
            invoice = BillingDocument.objects.create(tenant=self.org, organization=self.org, customer=customer, document_type='invoice', number=f'INV-{index}', issue_date=date(2026, 10, 1), currency=currency, created_by=self.admin)
            line = BillingLineItem.objects.create(tenant=self.org, organization=self.org, document=invoice, product=self.product, description='Historical router label', quantity=2, unit_price=20)
            sale = InventorySale.objects.create(tenant=self.org, invoice=invoice, stock_deducted=True, completed_at=datetime(2026, 10, 1, 22, 30, tzinfo=dt_timezone.utc))
            InventorySaleLine.objects.create(tenant=self.org, sale=sale, billing_line=line, net_revenue=40, cost_total=20)
        Product.objects.filter(tenant=self.org, pk=self.product.pk).update(name='Renamed router', buying_price=999)
        with timezone.override('Africa/Dar_es_Salaam'):
            params = {'date_from': '2026-10-02', 'date_to': '2026-10-02'}
            sales = self.client.get(self.url('sales'), params)
            self.assertEqual(sales.context['result_count'], 2)
            self.assertEqual(dict(sales.context['summaries']), {'Net revenue (TZS)': Decimal('40'), 'Net revenue (USD)': Decimal('40')})
            self.assertContains(sales, 'Historical router label')
            profit = self.client.get(self.url('gross-profit'), params)
            totals = dict(profit.context['summaries'])
            self.assertEqual(totals['Stock cost (TZS)'], Decimal('40'))
            self.assertEqual(totals['Gross profit on TZS sales (TZS)'], Decimal('20'))
            self.assertIsNone([row for row in profit.context['rows'] if row[-1] == 'USD'][0][7])
            self.assertEqual(self.client.get(self.url('sales'), {'date_to': '2026-10-01'}).context['result_count'], 0)

    def test_historical_reports_are_separate_from_live_purchases(self):
        job = ImportJob.objects.create(tenant=self.org, import_type='historical_purchases', created_by=self.admin)
        for kind in ('purchase', 'sale'):
            HistoricalInventoryRecord.objects.create(tenant=self.org, record_type=kind, record_date=date(2026, 10, 1), sku='HIST', description='Archived', quantity=2, unit_amount=15, total_amount=30, import_job=job)
            response = self.client.get(self.url('historical-purchases' if kind == 'purchase' else 'historical-sales'), {'date_from': '2026-10-01', 'date_to': '2026-10-01'})
            self.assertEqual(response.context['result_count'], 1)
            self.assertEqual(response.context['summaries'][0][1], 30)
        self.assertEqual(self.client.get(self.url()).context['summaries'][0][1], 0)
        self.assertFalse(InventoryBalance.objects.filter(tenant=self.org).exists())

    def test_rankings_thresholds_and_serial_report(self):
        InventorySettings.objects.create(tenant=self.org, fast_moving_min_units=5, slow_moving_max_units=2)
        InventoryBalance.objects.create(tenant=self.org, product=self.product, quantity=1, average_cost=10)
        self.product.reorder_threshold = 2
        self.product.save()
        self.assertEqual(len(report_rows(self.org, 'low-stock')[1]), 1)
        for _ in range(2):
            StockMovement.objects.create(tenant=self.org, product=self.product, movement_type='sale_out', quantity=-3, balance_after=1, unit_cost=10)
        for name in ('best-selling', 'fast-moving'):
            self.assertEqual(report_rows(self.org, name)[1][0][2], 6)
        for name in ('slow-moving', 'dead-stock'):
            self.assertEqual(report_rows(self.org, name)[1], [])
        StockUnit.objects.create(tenant=self.org, product=self.product, serial_number='REPORT-SERIAL')
        self.assertEqual(report_rows(self.org, 'serialized')[1][0][2], 'REPORT-SERIAL')
